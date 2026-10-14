# -*- coding: utf-8 -*-
from openpyxl import load_workbook
files = [
 r"D:/WK/N CRT/北部山头3940平台2026/7、单日数据/北部山头3940平台-1008/3940平台 1008炮孔数据报告.xlsx",
 r"D:/WK/N CRT/北部山头3940平台2026/3、北部山头3940平台炮孔数据报告-2026综合.xlsx",
]
for f in files:
    print("="*100); print("FILE:", f)
    wb = load_workbook(f, data_only=False)
    print("  sheets:", wb.sheetnames)
    for sn in wb.sheetnames:
        ws = wb[sn]
        print("  --- sheet [%s] dims=%s max_row=%d max_col=%d" % (sn, ws.dimensions, ws.max_row, ws.max_column))
        print("      merged:", [str(m) for m in ws.merged_cells.ranges])
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if i > 60: 
                print("      ...(truncated)"); break
            print("      r%02d|" % i, " | ".join(repr(v) for v in row))
    wb.close()
