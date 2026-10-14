# -*- coding: utf-8 -*-
import sys; sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook
p = r"D:/WK/N CRT/_probe/sb3960/北部山头3960平台2026/3、北部山头3960平台炮孔数据报告-2026综合.xlsx"
wb = load_workbook(p); ws = wb.active
print("dims", ws.dimensions, "merged", [str(m) for m in ws.merged_cells.ranges][:12])
rows = list(ws.iter_rows(values_only=True))
print("总行数", len(rows))
for i in range(len(rows)-14, len(rows)):
    print("  r%03d|" % (i+1), " | ".join(repr(v) for v in rows[i][:6]))
wb.close()
