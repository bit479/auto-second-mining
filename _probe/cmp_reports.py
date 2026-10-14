# -*- coding: utf-8 -*-
from openpyxl import load_workbook
def dump(f, label):
    print("="*90); print(label, f)
    wb = load_workbook(f, data_only=False)
    for sn in wb.sheetnames:
        ws = wb[sn]
        print(" [%s] max_row=%d max_col=%d merged=%s" % (sn, ws.max_row, ws.max_column, [str(m) for m in ws.merged_cells.ranges]))
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            print("  r%02d|" % i, " | ".join(repr(v) for v in row))
    wb.close()
dump(r"D:/WK/N CRT/北部山头3940平台2026/7、单日数据/北部山头3940平台-0915/3940平台 0915炮孔数据报告.xlsx","0915单日")
dump(r"D:/WK/N CRT/北部山头3940平台2026/7、单日数据/北部山头3940平台-0922/3940平台 0922炮孔数据报告.xlsx","0922单日")
