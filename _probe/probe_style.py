# -*- coding: utf-8 -*-
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
def sty(f, label, rows):
    print("="*90); print(label)
    wb = load_workbook(f); ws = wb[wb.sheetnames[0]]
    for r in rows:
        for c in range(1, 7):
            cell = ws.cell(r, c)
            print("  %s%d value=%r font=(%s,%s,bold=%s) align=(%s,%s,wrap=%s) numfmt=%r border=(%s,%s,%s,%s)" % (
                get_column_letter(c), r, cell.value, cell.font.name, cell.font.size, cell.font.bold,
                cell.alignment.horizontal, cell.alignment.vertical, cell.alignment.wrap_text,
                cell.number_format, cell.border.left.style, cell.border.right.style, cell.border.top.style, cell.border.bottom.style))
    print("  col widths:", {k: v.width for k,v in ws.column_dimensions.items()})
    print("  row heights:", {k: v.height for k,v in ws.row_dimensions.items()})
    wb.close()
sty(r"D:/WK/N CRT/北部山头3940平台2026/7、单日数据/北部山头3940平台-1008/3940平台 1008炮孔数据报告.xlsx","1008单日报告",[1,2,3,5,8])
sty(r"D:/WK/N CRT/北部山头3940平台2026/3、北部山头3940平台炮孔数据报告-2026综合.xlsx","综合总表",[1,2,3,6,13,15])
