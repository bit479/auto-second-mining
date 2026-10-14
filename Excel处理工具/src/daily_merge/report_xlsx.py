"""炮孔数据报告 .xlsx（总表）的并入。

总表是"从上到下依次堆叠的日报告块"：
每一块 = 标题行 + 表头行 + 各品位档明细 + 小计 + 合计。

并入时取单日报告的**汇总块**（标题行 → 合计行），在总表末尾空一行后整块粘贴，
块内的公式按行偏移整体平移（``=F5/D5*100`` → ``=F27/D27*100``），
字体、边框、对齐、数字格式、合并单元格一并带走。
"""

from __future__ import annotations

import copy as _copy
import os
from typing import Any

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator

from .schema import (KINDS, REPORT_SUBTOTAL_TOKENS, REPORT_TOTAL_TOKEN,
                     REPORT_TOTAL_TOKENS, REPORT_TYPE_COLUMN)

__all__ = ["find_summary_block", "existing_titles", "check_totals", "merge_report"]

_BLANK_SEPARATOR = 1


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _last_content_row(ws) -> int:
    last = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and str(cell.value).strip() != "":
                last = max(last, cell.row)
    return last


def _block_columns(ws, r1: int, r2: int) -> int:
    """块内实际用到的最大列号。"""
    ncol = 1
    for row in ws.iter_rows(min_row=r1, max_row=r2):
        for cell in row:
            if cell.value is not None and str(cell.value).strip() != "":
                ncol = max(ncol, cell.column)
    return ncol


def find_summary_block(ws) -> tuple[int, int]:
    """定位汇总块的首尾行（含标题行与合计行）。找不到抛 ValueError。"""
    first = None
    total = None
    for row in ws.iter_rows(min_col=1, max_col=1):
        text = _text(row[0].value)
        if first is None and text:
            first = row[0].row
        if text in REPORT_TOTAL_TOKENS:
            total = row[0].row
            break
    if total is None:
        raise ValueError("单日报告里没找到「%s」行，无法确定汇总块范围"
                         % "/".join(REPORT_TOTAL_TOKENS))
    if first is None:
        first = 1
    return first, total


def existing_titles(ws) -> list[str]:
    """总表里已有的日报告块标题（用于防止同一日重复并入）。"""
    titles = []
    for row in ws.iter_rows(min_col=1, max_col=1):
        value = row[0].value
        if isinstance(value, str) and "炮孔数据报告" in value:
            titles.append(_text(value))
    return titles


# 合计行所在列：体积 / 重量 / 金属量（平均品位是推导列，不参与求和校验）
_SUM_COLUMNS = ((3, "体积"), (4, "重量"), (6, "金属量"))


def check_totals(src_path: str, tolerance: float = 0.05) -> dict:
    """用文件里缓存的计算值，复核"合计 = 各明细行之和"。

    只做比对并报告差异，绝不改动源数据 —— 源报告本身写错公式时，
    这里会明确提示，避免错误被原样搬进总表。
    """
    wb = load_workbook(src_path, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        try:
            r1, r2 = find_summary_block(ws)
        except ValueError as exc:
            return {"checked": False, "reason": str(exc)}

        sums = {name: 0.0 for _col, name in _SUM_COLUMNS}
        details = 0
        for r in range(r1 + 1, r2):
            label = _text(ws.cell(r, 1).value)
            if not label or label in REPORT_TOTAL_TOKENS or label in REPORT_SUBTOTAL_TOKENS:
                continue
            details += 1
            for col, name in _SUM_COLUMNS:
                value = ws.cell(r, col).value
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    sums[name] += float(value)

        total = {}
        for col, name in _SUM_COLUMNS:
            value = ws.cell(r2, col).value
            total[name] = float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else None

        mismatches = []
        for _col, name in _SUM_COLUMNS:
            got = total[name]
            if got is None:
                continue
            if abs(got - sums[name]) > tolerance:
                mismatches.append({"column": name, "detail_sum": round(sums[name], 3),
                                   "total": round(got, 3),
                                   "diff": round(got - sums[name], 3)})
        return {
            "checked": True,
            "detail_rows": details,
            "total_label": _text(ws.cell(r2, 1).value),
            "detail_sum": {k: round(v, 3) for k, v in sums.items()},
            "total": {k: (round(v, 3) if v is not None else None) for k, v in total.items()},
            "mismatches": mismatches,
        }
    finally:
        wb.close()


def _copy_style(src_cell, dst_cell) -> None:
    dst_cell.font = _copy.copy(src_cell.font)
    dst_cell.border = _copy.copy(src_cell.border)
    dst_cell.fill = _copy.copy(src_cell.fill)
    dst_cell.alignment = _copy.copy(src_cell.alignment)
    dst_cell.protection = _copy.copy(src_cell.protection)
    dst_cell.number_format = src_cell.number_format


def _translate(value: Any, origin: str, target: str) -> Any:
    if isinstance(value, str) and value.startswith("="):
        try:
            return Translator(value, origin=origin).translate_formula(target)
        except Exception:  # noqa: BLE001 - 平移失败就原样保留
            return value
    return value


def merge_report(src_path: str, dst_path: str, *,
                 with_appendix: bool = False,
                 separator: int = _BLANK_SEPARATOR,
                 skip_duplicates: bool = True,
                 dry_run: bool = False) -> dict:
    """把 ``src_path`` 单日报告的汇总块追加到 ``dst_path`` 总表末尾。"""
    src_wb = load_workbook(src_path, data_only=False)
    src_ws = src_wb[src_wb.sheetnames[0]]

    r1, r2 = find_summary_block(src_ws)
    if with_appendix:
        r2 = max(r2, _last_content_row(src_ws))
    ncol = _block_columns(src_ws, r1, r2)
    title = _text(src_ws.cell(r1, 1).value)

    dst_wb = load_workbook(dst_path, data_only=False)
    dst_ws = dst_wb[dst_wb.sheetnames[0]]
    dst_last = _last_content_row(dst_ws)

    if skip_duplicates and title and title in existing_titles(dst_ws):
        return {
            "kind": KINDS[2],
            "src": os.path.abspath(src_path),
            "dst": os.path.abspath(dst_path),
            "written": False,
            "skipped": True,
            "title": title,
            "block_rows": r2 - r1 + 1,
            "columns": ncol,
            "formulas": 0,
            "insert_at": 0,
            "before_rows": dst_last,
            "after_rows": dst_last,
            "added": 0,
            "reason": "总表中已存在「%s」，未重复并入" % title,
        }

    start = dst_last + 1 + max(separator, 0)
    delta = start - r1

    formulas = sum(
        1
        for row in src_ws.iter_rows(min_row=r1, max_row=r2, max_col=ncol)
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    )

    if not dry_run:
        for row in src_ws.iter_rows(min_row=r1, max_row=r2, max_col=ncol):
            for src_cell in row:
                if src_cell.value is None:
                    continue
                dst_cell = dst_ws.cell(src_cell.row + delta, src_cell.column)
                dst_cell.value = _translate(src_cell.value, src_cell.coordinate, dst_cell.coordinate)
                _copy_style(src_cell, dst_cell)
        # 合并单元格（只搬块内的）
        for rng in src_ws.merged_cells.ranges:
            if rng.min_row >= r1 and rng.max_row <= r2 and rng.max_col <= ncol:
                dst_ws.merge_cells(
                    start_row=rng.min_row + delta, start_column=rng.min_col,
                    end_row=rng.max_row + delta, end_column=rng.max_col,
                )
        # 行高
        for r in range(r1, r2 + 1):
            height = src_ws.row_dimensions[r].height
            if height:
                dst_ws.row_dimensions[r + delta].height = height

        dst_wb.save(dst_path)

    return {
        "kind": KINDS[2],
        "src": os.path.abspath(src_path),
        "dst": os.path.abspath(dst_path),
        "written": not dry_run,
        "skipped": False,
        "title": title,
        "block_rows": r2 - r1 + 1,
        "columns": ncol,
        "formulas": formulas,
        "insert_at": start,
        "before_rows": dst_last,
        "after_rows": dst_last + (0 if dry_run else (separator + r2 - r1 + 1)),
        "added": r2 - r1 + 1,
    }
