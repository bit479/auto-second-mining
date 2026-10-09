"""Excel 读取层。

能力清单：
* 支持 ``.xlsx / .xlsm / .xls / .csv / .tsv``；
* 自动定位表头行（跳过合并的大标题、说明行、空行）；
* 自动丢弃"小计/合计/总计"等汇总行（可关闭）；
* 单元格值清洗（去空格、全角转半角、数值转 float）；
* 一次读出全部 sheet，或按名字/序号指定。
"""

from __future__ import annotations

import csv
import glob
import os
import re
from typing import Any, Iterable, Sequence

from openpyxl import load_workbook

from .model import DataTable, is_blank, to_number, to_text

__all__ = [
    "read_workbook",
    "read_table",
    "read_many",
    "find_input_files",
    "detect_header_row",
    "SUPPORTED_EXT",
]

SUPPORTED_EXT = (".xlsx", ".xlsm", ".xltx", ".xltm", ".xls", ".csv", ".tsv")

# 这些行不是数据行，解析业务数据时默认剔除
SUMMARY_TOKENS = ("小计", "合计", "总计", "共计", "sum", "total", "subtotal")


def find_input_files(paths: Iterable[str] | str | None = None,
                     directory: str | None = None,
                     pattern: str = "*.xls*") -> list[str]:
    """把各种输入形式解析成文件列表。

    ``paths`` 可以是单个路径、路径列表、或含通配符的字符串；
    ``directory`` + ``pattern`` 用于按目录批量收集。
    会自动跳过 Excel 临时文件（``~$`` 开头）。
    """
    found: list[str] = []
    if paths:
        if isinstance(paths, str):
            paths = [paths]
        for p in paths:
            if any(ch in p for ch in "*?"):
                found.extend(sorted(glob.glob(p)))
            elif os.path.isdir(p):
                found.extend(sorted(glob.glob(os.path.join(p, pattern))))
            else:
                found.append(p)
    if directory:
        found.extend(sorted(glob.glob(os.path.join(directory, pattern))))

    result, seen = [], set()
    for f in found:
        if not os.path.isfile(f):
            continue
        if os.path.basename(f).startswith("~$"):  # Excel 打开时产生的临时文件
            continue
        if os.path.splitext(f)[1].lower() not in SUPPORTED_EXT:
            continue
        key = os.path.abspath(f)
        if key not in seen:
            seen.add(key)
            result.append(f)
    return result


def detect_header_row(matrix: Sequence[Sequence[Any]], scan: int = 12) -> int:
    """在原始矩阵里猜表头行号（0-based）。

    判据：该行非空文本单元格最多，且紧随其后的行里数值单元格最多。
    """
    limit = min(scan, len(matrix) - 1) if len(matrix) > 1 else len(matrix)
    best_row, best_score = 0, -1.0
    for i in range(max(limit, 1)):
        row = matrix[i]
        n_text = sum(1 for v in row if not is_blank(v) and to_number(v) is None)
        n_all = sum(1 for v in row if not is_blank(v))
        if n_all == 0:
            continue
        score = n_text / n_all * 2 + n_text * 0.1
        nxt = matrix[i + 1] if i + 1 < len(matrix) else []
        score += sum(1 for v in nxt if to_number(v) is not None) * 0.3
        if score > best_score:
            best_row, best_score = i, score
    return best_row


def _matrix_from_worksheet(ws, max_rows: int | None = None, max_cols: int | None = None) -> list[list[Any]]:
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append(list(r))
        if max_rows and len(rows) >= max_rows:
            break
    if max_cols:
        rows = [r[:max_cols] for r in rows]
    # 去掉尾部全空行
    while rows and all(is_blank(v) for v in rows[-1]):
        rows.pop()
    return rows


def _read_xlsx(path: str, sheet: str | int | None = None) -> list[tuple[str, list[list[Any]]]]:
    wb = load_workbook(path, data_only=True, read_only=False)
    try:
        if sheet is None:
            names = wb.sheetnames
        elif isinstance(sheet, int):
            names = [wb.sheetnames[sheet]]
        else:
            names = [sheet]
        return [(n, _matrix_from_worksheet(wb[n])) for n in names]
    finally:
        wb.close()


def _read_xls(path: str, sheet: str | int | None = None) -> list[tuple[str, list[list[Any]]]]:
    try:
        import xlrd  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("读取 .xls 需要安装 xlrd：pip install xlrd") from exc
    book = xlrd.open_workbook(path)
    names = book.sheet_names() if sheet is None else (
        [book.sheet_names()[sheet]] if isinstance(sheet, int) else [sheet]
    )
    out = []
    for n in names:
        sh = book.sheet_by_name(n)
        out.append((n, [sh.row_values(i) for i in range(sh.nrows)]))
    return out


def _read_csv(path: str, delimiter: str = ",") -> list[tuple[str, list[list[Any]]]]:
    encoding = "utf-8-sig"
    for enc in ("utf-8-sig", "gbk", "utf-8"):
        try:
            with open(path, "r", encoding=enc, newline="") as fh:
                list(csv.reader(fh, delimiter=delimiter))
            encoding = enc
            break
        except UnicodeDecodeError:
            continue
    with open(path, "r", encoding=encoding, newline="") as fh:
        rows = [row for row in csv.reader(fh, delimiter=delimiter)]
    return [("Sheet1", rows)]


def read_workbook(path: str, sheet: str | int | None = None) -> list[tuple[str, list[list[Any]]]]:
    """读出文件中所有（或指定）sheet 的原始矩阵。"""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        return _read_xlsx(path, sheet)
    if ext == ".xls":
        return _read_xls(path, sheet)
    if ext in (".csv", ".tsv"):
        return _read_csv(path, "\t" if ext == ".tsv" else ",")
    raise ValueError("不支持的文件类型：%s" % path)


def read_table(path: str,
               sheet: str | int | None = None,
               header_row: int | None = None,
               numeric_columns: Sequence[str] | None = None,
               drop_summary_rows: bool = True,
               summary_column: str | None = None) -> list[DataTable]:
    """把一个文件读成一张或多张 DataTable。

    Parameters
    ----------
    header_row
        表头所在行（0-based）；不传则自动探测。
    numeric_columns
        需要强制转成 float 的列名。
    drop_summary_rows
        是否丢弃"小计/合计"这类汇总行，默认丢弃。
    summary_column
        在哪个列上判断汇总行；不传则取第一列。
    """
    tables: list[DataTable] = []
    for name, matrix in read_workbook(path, sheet):
        if not matrix:
            continue
        hr = header_row if header_row is not None else detect_header_row(matrix)
        table = DataTable.from_matrix(matrix[hr:], name=name, source=os.path.basename(path))
        if not table.columns:
            continue
        if drop_summary_rows:
            col = summary_column or table.columns[0]
            table = table.drop_rows_matching(col, SUMMARY_TOKENS)
        if numeric_columns:
            table = table.cast({c: to_number for c in numeric_columns if c in table.columns})
        tables.append(table)
    return tables


def read_many(paths: Iterable[str],
              sheet: str | int | None = None,
              header_row: int | None = None,
              numeric_columns: Sequence[str] | None = None,
              drop_summary_rows: bool = True,
              summary_column: str | None = None) -> list[DataTable]:
    """批量读取，自动跳过打不开的文件并保留其余结果。"""
    out: list[DataTable] = []
    for p in paths:
        try:
            for t in read_table(p, sheet, header_row, numeric_columns, drop_summary_rows, summary_column):
                t.source = os.path.basename(p)
                out.append(t)
        except Exception as exc:  # noqa: BLE001 - 单文件失败不影响整体
            print("[警告] 跳过无法读取的文件 %s：%s" % (p, exc))
    return out


def natural_key_compat(text: str) -> list:
    """自然排序键：让"工作簿2"排在"工作簿10"前面。"""
    return [int(s) if s.isdigit() else s.lower()
            for s in re.split(r"(\d+)", os.path.basename(str(text)))]


def column_letter_width(values: Iterable[Any]) -> float:
    """按内容估算列宽（中文按 2 个字符宽算）。"""
    widest = 8.0
    for v in values:
        text = to_text(v)
        if not text:
            continue
        width = sum(2.0 if ord(ch) > 127 else 1.0 for ch in text)
        widest = max(widest, width + 2)
    return min(widest, 40.0)


_UNSAFE = re.compile(r"[\\/:*?\[\]]")


def safe_sheet_title(name: str, fallback: str = "Sheet") -> str:
    """Excel 工作表名不能含 \\ / : * ? [ ] 且不超过 31 字符。"""
    clean = _UNSAFE.sub("_", to_text(name))[:31]
    return clean or fallback
