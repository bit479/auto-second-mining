"""自检：用 handle 目录下的 3 个工作簿跑汇总，并与人工成果逐项比对。"""

from __future__ import annotations

import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from src.blast_report.builder import MergeOptions, merge_reports  # noqa: E402
from src.excel_core.reader import find_input_files, natural_key_compat  # noqa: E402

SRC_DIR = r"D:\WK\N CRT\handle"
OUT_DIR = os.path.join(ROOT, "_out")
OUTPUT = os.path.join(OUT_DIR, "3940平台 1008炮孔数据报告.xlsx")

# 人工成果（3940平台 1008炮孔数据报告.xlsx，公式已算出的值）
EXPECTED = [
    ("1.500-3.000", "1", 515.29, 1391.282, 2.16, 30.052),
    ("小计",        "",  515.29, 1391.282, 2.160022, 30.052),
    ("1.000-1.500", "2", 445.042, 1201.614, 1.44, 17.303),
    ("小计",        "",  445.042, 1201.614, 1.44, 17.303),
    ("0.500-1.000", "1", 1229.228, 3318.915, 0.65, 21.557),
    ("0.500-1.000", "1", 714.095, 1928.056, 0.68, 13.111),
    ("小计",        "",  1943.323, 5246.971, 0.660724, 34.668),
    ("合计",        "",  2903.655, 7839.867, 1.046229, 82.023),
]


def main() -> int:
    files = find_input_files(directory=SRC_DIR, pattern="工作簿*.xlsx")
    files.sort(key=natural_key_compat)
    print("输入：", [os.path.basename(f) for f in files])

    result = merge_reports(files, OUTPUT, MergeOptions())
    print("输出：", result["output"])
    print("品位档顺序：", result["tiers"])

    # 用 --values 模式再生成一份纯数值文件，便于程序化比对
    from openpyxl import load_workbook

    values_out = os.path.join(OUT_DIR, "_values.xlsx")
    merge_reports(files, values_out, MergeOptions(use_formula=False, with_appendix=False))
    ws = load_workbook(values_out, data_only=True).active

    ok = True
    print("\n%-14s %-5s %12s %12s %10s %10s   %s"
          % ("类型", "体号", "体积", "重量", "品位", "金属量", "比对"))
    for i, exp in enumerate(EXPECTED):
        r = 3 + i
        got = [ws.cell(r, c).value for c in range(1, 7)]
        diffs = []
        for label, e, g in zip(["体积", "重量", "品位", "金属量"], exp[2:], got[2:]):
            tol = 1e-3 if label == "品位" else 1e-3
            if g is None or abs(float(g) - float(e)) > tol:
                diffs.append("%s 期望 %s 实得 %s" % (label, e, g))
        text_ok = str(got[0]) == exp[0] and str(got[1] or "") == exp[1]
        if not text_ok:
            diffs.append("文本 期望 %r 实得 %r" % (exp[:2], got[:2]))
        flag = "OK" if not diffs else "× " + "; ".join(diffs)
        ok = ok and not diffs
        print("%-14s %-5s %12s %12s %10s %10s   %s"
              % (got[0], got[1] or "", got[2], got[3], got[4], got[5], flag))

    print("\n结论：", "全部一致 ✓" if ok else "存在差异 ✗")
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
