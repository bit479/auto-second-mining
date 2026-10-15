"""自检：**新口径**（按每个文件的「总计」行汇总）。

用 3940 平台 1011 的 4 个 `工作表*.xls` 跑一遍，把汇总区逐格与人工成果比对：

* 每份文件只取「总计」行 → 一个矿块；
* 档位由矿块品位反推（工作表3 的总计跨了两个源档位，只能这么定）；
* 档位从高到低、同档内保持文件顺序；
* 体号在最终顺序上重新编号 1→N。

用法::

    python tests/test_by_total.py
"""

from __future__ import annotations

import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from openpyxl import load_workbook  # noqa: E402

from src.blast_report import naming  # noqa: E402
from src.blast_report.builder import MergeOptions, merge_reports  # noqa: E402
from src.blast_report.schema import band_of_grade  # noqa: E402

SRC = r"C:\Users\Administrator\Desktop\3940平台\1011"
DRAFT = os.path.join(SRC, "3940平台 1011炮孔数据报告.xlsx")
OUTPUT = os.path.join(ROOT, "_out", "1011炮孔数据报告.xlsx")

INPUTS = [os.path.join(SRC, "工作表%d.xls" % i) for i in range(1, 5)]

# 人工成果的汇总区（r3 ~ r9）：值为数值或公式字符串
EXPECTED = [
    ("1.000-1.500", "1", 1141.37, 3081.698, 1.404, 43.267),
    ("小计", "", "=C3", "=D3", "=E3", "=F3"),
    ("0.500-1.000", "2", 333.777, 901.198, 0.557, 5.021),
    ("0.500-1.000", "3", 423.117, 1142.415, 0.772, 8.817),
    ("0.500-1.000", "4", 1329.697, 3572.909, 0.658, 23.523),
    ("小计", "", "=C5+C6+C7", "=D5+D6+D7", "=F8/D8*100", "=F5+F6+F7"),
    ("合计", "", "=C4+C8", "=D4+D8", "=F9/D9*100", "=F4+F8"),
]

_checks: list[tuple[bool, str]] = []


def check(ok: bool, message: str) -> None:
    _checks.append((bool(ok), message))
    print("  %s  %s" % ("OK  " if ok else "FAIL", message))


def _same(a, b) -> bool:
    """数值按容差比；空值与空串等价；其余按字符串比。"""
    if a is None and b == "":
        return True
    if b is None and a == "":
        return True
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(float(a) - float(b)) < 1e-6
    return a == b


def _body(value):
    """体号列读回来可能是 int / '1' / None，统一成字符串比较。"""
    if value is None or value == "":
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def test_band_of_grade() -> None:
    print("\n[档位反推]")
    cases = [
        (0.0, "0.000-0.500"), (0.499, "0.000-0.500"),
        (0.5, "0.500-1.000"),                       # 边界：下限含
        (0.557, "0.500-1.000"), (0.772, "0.500-1.000"), (0.658, "0.500-1.000"),
        (0.999, "0.500-1.000"),
        (1.0, "1.000-1.500"), (1.404, "1.000-1.500"), (1.499, "1.000-1.500"),
        (1.5, "1.500-3.000"), (2.999, "1.500-3.000"),
        (3.0, "3.000-999.000"), (12.5, "3.000-999.000"),
    ]
    bad = [(v, band_of_grade(v), want) for v, want in cases if band_of_grade(v) != want]
    check(not bad, "14 个品位值都落到正确档位" + ("" if not bad else " → %s" % bad))


def test_naming() -> None:
    print("\n[标题 / 默认输出名]")
    title = naming.guess_title(INPUTS, DRAFT)
    check(title == "北部山头3940平台1011炮孔数据报告", "标题推断 = %r" % title)

    default_out = naming.default_output_path(INPUTS, title)
    check(os.path.basename(default_out) == "北部山头3940平台1011炮孔数据报告.xlsx",
          "默认输出文件名 = %r" % os.path.basename(default_out))
    check(os.path.dirname(default_out) == SRC, "默认落在输入文件所在目录")

    # 路径里已经带矿名时不要重复补前缀
    deep = r"D:\WK\N CRT\北部山头3940平台2026\7、单日数据\北部山头3940平台-1008\工作簿1.xlsx"
    check(naming.guess_title([deep], None) == "北部山头3940平台1008炮孔数据报告",
          "带矿名的路径 → %r" % naming.guess_title([deep], None))


def test_merge() -> int:
    print("\n[汇总]")
    result = merge_reports(INPUTS, OUTPUT, MergeOptions())
    check(result["mode"] == "total", "默认方式 = %s" % result["mode"])
    check(result["blocks"] == 4, "产出 4 个矿块（4 份文件各取一行总计）")
    check(result["tiers"] == ["1.000-1.500", "0.500-1.000"],
          "档位从高到低：%s" % result["tiers"])
    check(result["title"] == "北部山头3940平台1011炮孔数据报告",
          "标题 = %r" % result["title"])

    ws = load_workbook(OUTPUT, data_only=False).active
    check(ws.cell(2, 1).value == "类型" and ws.cell(2, 6).value == "金属量（百克）",
          "表头 6 列正确")

    print("\n  %-4s %-13s %-5s %-11s %-11s %-11s %-11s  %s"
          % ("行", "类型", "体号", "体积", "重量", "品位", "金属量", "比对"))
    ok = True
    for i, exp in enumerate(EXPECTED):
        row = 3 + i
        got = [ws.cell(row, c).value for c in range(1, 7)]
        diffs = []
        for c in range(6):
            a, b = exp[c], got[c]
            if c == 1:
                a, b = _body(a), _body(b)
            if not _same(a, b):
                diffs.append("第%d列 %r≠%r" % (c + 1, exp[c], got[c]))
        ok = ok and not diffs
        print("  %-4d %-13s %-5s %-11s %-11s %-11s %-11s  %s"
              % (row, got[0], got[1] or "", got[2], got[3], got[4], got[5],
                 "OK" if not diffs else "× " + "; ".join(diffs)))
    check(ok, "汇总区 r3~r9 逐格与人工成果一致（含公式写法）")
    return 0 if ok else 1


def test_appendix() -> None:
    print("\n[附录]")
    ws = load_workbook(OUTPUT, data_only=False).active
    titles = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)
              if ws.cell(r, 1).value == "3DMine 炮孔数据库报告"]
    check(len(titles) == 4, "下方附了 4 份原始报告")

    # 工作表4 的「总计」下面挂着一行重复数据，必须被截掉
    stale = [r for r in range(1, ws.max_row + 1)
             if ws.cell(r, 1).value == "1.000-1.500" and ws.cell(r, 3).value == 1141.37]
    check(len(stale) == 1, "总计行下方的杂项行未进附录（只保留汇总区的 1 处）")


def test_draft_compare() -> None:
    print("\n[与桌面上的草稿文件直接比对]")
    if not os.path.isfile(DRAFT):
        check(True, "草稿文件不存在，跳过（不影响结论）")
        return
    ref = load_workbook(DRAFT, data_only=False).active
    new = load_workbook(OUTPUT, data_only=False).active
    diffs = []
    for row in range(1, 10):
        for col in range(1, 7):
            a, b = ref.cell(row, col).value, new.cell(row, col).value
            if not _same(a, b):
                diffs.append("r%dC%d %r≠%r" % (row, col, a, b))
    check(not diffs, "r1~r9 与草稿逐格一致" + ("" if not diffs else " → %s" % diffs))


def main() -> int:
    print("=" * 78)
    for path in INPUTS:
        if not os.path.isfile(path):
            print("[跳过] 找不到输入文件：%s" % path)
            return 0
    print("输入：", [os.path.basename(p) for p in INPUTS])
    print("-" * 78)

    test_band_of_grade()
    test_naming()
    failed_merge = test_merge()
    test_appendix()
    test_draft_compare()

    failed = [msg for ok, msg in _checks if not ok]
    print("\n" + "=" * 78)
    print("共 %d 项检查，%d 项通过" % (len(_checks), len(_checks) - len(failed)))
    if failed:
        for msg in failed:
            print("  FAIL  %s" % msg)
        print("结论：存在差异 ✗")
        return 1
    print("结论：全部一致 ✓")
    return 0 if not failed_merge else 1


if __name__ == "__main__":
    raise SystemExit(main())
