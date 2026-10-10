# -*- coding: utf-8 -*-
"""三件套生成：报告 xlsx / 矿界图 DXF / 矿块边界 .3ds

绘图细则来自人工图实测：
  十字 "+"    ±0.5m 两条线（层 点）
  轨迹线       孔口向西 1.2m 水平线（层 钻孔）
  工程号       孔口偏移 (-2.28, +0.40)，h=0.4，颜色 138（层 钻孔）
  品位         孔口偏移 (-3.19, -0.35)，h=0.7，按档着色（层 钻孔）
  1m 方格      按矿界裁剪后的 1×1 闭合小多边形（层 爆区方格网）
  方格品位     方格中心，h=0.1333，按档着色
  矿块号       "n号矿块" h=3.0 + 斜向引线（从本块边界出发）
"""
from __future__ import annotations

import math
from pathlib import Path

from shapely.geometry import LineString, Point, Polygon, box
from shapely.ops import unary_union

LEVEL_COLOR = {"L1": 4, "L2": 5, "L3": 1, "L4": 2}          # AutoCAD 色号（青/蓝/红/黄）
LEVEL_RGB = {"L1": (0, 255, 255), "L2": (0, 0, 255),
             "L3": (255, 0, 0), "L4": (255, 255, 0)}
HID_COLOR = 138          # 工程号颜色（人工图实测）


def level_id_of(grade: float):
    if grade >= 3.0:
        return "L4"
    if grade >= 1.5:
        return "L3"
    if grade >= 1.0:
        return "L2"
    if grade >= 0.5:
        return "L1"
    return None


def grid_cells(poly: Polygon, size: float = 1.0):
    """按矿界裁剪的 1×1 m 方格（闭合小多边形）。"""
    out = []
    minx, miny, maxx, maxy = poly.bounds
    gx0 = math.floor(minx / size) * size
    gy0 = math.floor(miny / size) * size
    x = gx0
    while x < maxx:
        y = gy0
        while y < maxy:
            c = box(x, y, x + size, y + size)
            inter = c.intersection(poly)
            if not inter.is_empty and inter.area > 1e-9:
                if inter.geom_type == "Polygon":
                    out.append(inter)
                else:
                    for g in getattr(inter, "geoms", []):
                        if g.geom_type == "Polygon" and g.area > 1e-9:
                            out.append(g)
            y += size
        x += size
    return out


# ---------------------------------------------------------------- 报告 xlsx
def write_report_xlsx(blocks, meta, path: Path) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    title = "%s平台 %s炮孔数据报告" % (meta["platform"], meta["date"])
    ws["A1"] = title
    ws.merge_cells("A1:F1")
    ws["A1"].font = Font(bold=True, size=12)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    head = ["类型", "体号", "体积", "重量", "平均品位(Au)", "金属量（百克）"]
    for i, h in enumerate(head, 1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")
        c.fill = PatternFill("solid", fgColor="D9E2F3")

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    r = 3
    totals = [0.0, 0.0, 0.0]
    for lid in ("L4", "L3", "L2", "L1"):
        bs = [b for b in blocks if b.level_id == lid]
        if not bs:
            continue
        sv = st = sm = 0.0
        for b in bs:
            ws.cell(row=r, column=1, value=b.label)
            ws.cell(row=r, column=2, value=b.no)
            ws.cell(row=r, column=3, value=round(b.volume_m3, 3))
            ws.cell(row=r, column=4, value=round(b.tonnage_t, 3))
            ws.cell(row=r, column=5, value=round(b.grade, 3))
            ws.cell(row=r, column=6, value=round(b.metal_hg, 3))
            sv += b.volume_m3
            st += b.tonnage_t
            sm += b.metal_hg
            r += 1
        ws.cell(row=r, column=1, value="小计")
        ws.cell(row=r, column=3, value=round(sv, 3))
        ws.cell(row=r, column=4, value=round(st, 3))
        ws.cell(row=r, column=5, value=round(sm * 100.0 / st, 3) if st else 0)
        ws.cell(row=r, column=6, value=round(sm, 3))
        for c in range(1, 7):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="EDEDED")
        r += 1
        totals[0] += sv
        totals[1] += st
        totals[2] += sm

    ws.cell(row=r, column=1, value="合计")
    ws.cell(row=r, column=3, value=round(totals[0], 3))
    ws.cell(row=r, column=4, value=round(totals[1], 3))
    ws.cell(row=r, column=5, value=round(totals[2] * 100.0 / totals[1], 3) if totals[1] else 0)
    ws.cell(row=r, column=6, value=round(totals[2], 3))
    for c in range(1, 7):
        ws.cell(row=r, column=c).font = Font(bold=True)

    for row in ws.iter_rows(min_row=2, max_row=r, min_col=1, max_col=6):
        for cell in row:
            cell.border = border
            if cell.column >= 3:
                cell.number_format = "0.000"
    for col, w in zip("ABCDEF", (14, 8, 12, 12, 14, 14)):
        ws.column_dimensions[col].width = w
    wb.save(str(path))


# ---------------------------------------------------------------- DXF
def _nearest_hole_grade(cell: Polygon, holes):
    c = cell.representative_point()
    best, bd = None, 1e18
    for h in holes:
        d = (h.x - c.x) ** 2 + (h.y - c.y) ** 2
        if d < bd:
            bd, best = d, h
    return best


def write_dxf(blocks, holes_all, cells, meta, path: Path) -> None:
    from ezdxf import new

    doc = new("R2010")
    doc.header["$INSUNITS"] = 6
    for lid, name in (("L1", "0.500-1.000"), ("L2", "1.000-1.500"),
                      ("L3", "1.500-3.000"), ("L4", "3.000-999.000")):
        doc.layers.add(name, color=LEVEL_COLOR[lid])
    for n in ("爆区方格网", "钻孔", "点"):
        if n not in doc.layers:
            doc.layers.add(n, color=7)
    cn = doc.styles.add("CN_TTF", font="simsunb.ttf")
    cn.dxf.bigfont = ""
    msp = doc.modelspace()

    layer_of = {"L1": "0.500-1.000", "L2": "1.000-1.500",
                "L3": "1.500-3.000", "L4": "3.000-999.000"}
    # 1) 矿界（彩色闭合线）
    for b in blocks:
        if b.polygon is None:
            continue
        polys = list(b.polygon.geoms) if b.polygon.geom_type == "MultiPolygon" else [b.polygon]
        for p in polys:
            msp.add_lwpolyline([(x, y) for x, y in p.exterior.coords], close=True,
                               dxfattribs={"layer": layer_of[b.level_id],
                                           "color": LEVEL_COLOR[b.level_id],
                                           "lineweight": 30})

    # 2) 1m 方格 + 方格中心品位
    n_grid = 0
    for b in blocks:
        if b.polygon is None:
            continue
        color = LEVEL_COLOR[b.level_id]
        for cellp in grid_cells(b.polygon, 1.0):
            msp.add_lwpolyline([(x, y) for x, y in cellp.exterior.coords], close=True,
                               dxfattribs={"layer": "爆区方格网", "color": 7})
            n_grid += 1
            hh = _nearest_hole_grade(cellp, b.holes)
            if hh is not None:
                c = cellp.representative_point()
                msp.add_text("%.2f" % hh.grade_disp,
                             dxfattribs={"layer": "0", "color": color,
                                         "height": 0.1333, "style": "CN_TTF"},
                             ).set_placement((c.x - 0.16, c.y - 0.07))

    # 3) 炮孔符号 + 工程号 + 品位
    for h in holes_all:
        lid = level_id_of(h.grade)
        aci = LEVEL_COLOR.get(lid, 8)
        # 十字 ±0.5
        msp.add_line((h.x - 0.5, h.y), (h.x + 0.5, h.y),
                     dxfattribs={"layer": "点", "color": 7})
        msp.add_line((h.x, h.y - 0.5), (h.x, h.y + 0.5),
                     dxfattribs={"layer": "点", "color": 7})
        # 轨迹线：向西 1.2m
        msp.add_line((h.x, h.y), (h.x - 1.2, h.y),
                     dxfattribs={"layer": "钻孔", "color": 7, "lineweight": 35})
        # 工程号
        msp.add_text(h.hid, dxfattribs={"layer": "钻孔", "color": HID_COLOR,
                                        "height": 0.4, "style": "CN_TTF"},
                     ).set_placement((h.x - 2.28, h.y + 0.40))
        # 品位（分级色）
        msp.add_text("%.2f" % h.grade_disp,
                     dxfattribs={"layer": "钻孔", "color": aci,
                                 "height": 0.7, "style": "CN_TTF"},
                     ).set_placement((h.x - 3.19, h.y - 0.35))

    # 4) 矿块号：斜引线 + 文字（互不重叠）
    if blocks:
        allpts = [pt for b in blocks if b.polygon is not None
                  for pt in b.polygon.convex_hull.exterior.coords]
        gcx = sum(p[0] for p in allpts) / len(allpts)
        gcy = sum(p[1] for p in allpts) / len(allpts)
        placed = []
        for b in blocks:
            if b.polygon is None:
                continue
            cen = b.polygon.representative_point()
            dx, dy = cen.x - gcx, cen.y - gcy
            if abs(dx) < 1e-6 and abs(dy) < 1e-6:
                dx, dy = 1.0, 0.0
            L = math.hypot(dx, dy)
            ux, uy = dx / L, dy / L
            d = math.sqrt(b.polygon.area) * 0.9 + 8.0
            tx, ty = cen.x + ux * d, cen.y + uy * d
            # 简单避让
            for _ in range(60):
                clash = any(abs(tx - px) < 9.0 and abs(ty - py) < 4.0 for px, py in placed)
                if not clash:
                    break
                ty += 4.5
            placed.append((tx, ty))
            # 引线起点：本块边界上离文字最近的点
            bx, by = _nearest_on_boundary(b.polygon, tx, ty)
            msp.add_line((bx, by), (tx, ty),
                         dxfattribs={"layer": "0", "color": 7})
            msp.add_mtext("%d号矿块" % b.no,
                          dxfattribs={"layer": "0", "style": "CN_TTF",
                                      "char_height": 3.0}).set_location((tx, ty))

    # 5) 图名
    minx, miny, maxx, maxy = _bounds(holes_all)
    msp.add_text("%s平台 %s二次圈矿矿块图" % (meta["platform"], meta["date"]),
                 dxfattribs={"layer": "0", "color": 7, "height": 5.0,
                             "style": "CN_TTF"},
                 ).set_placement((minx, maxy + 12))
    doc.saveas(str(path))
    print("    方格 %d 个；炮孔 %d 个" % (n_grid, len(holes_all)))


def _bounds(holes):
    xs = [h.x for h in holes]
    ys = [h.y for h in holes]
    return min(xs), min(ys), max(xs), max(ys)


def _nearest_on_boundary(poly: Polygon, tx: float, ty: float):
    best = None
    bd = 1e18
    rings = []
    geoms = list(poly.geoms) if poly.geom_type == "MultiPolygon" else [poly]
    for g in geoms:
        rings.append(g.exterior)
        rings.extend(list(g.interiors))
    for ring in rings:
        for i in range(len(ring.coords) - 1):
            a = ring.coords[i]
            b = ring.coords[i + 1]
            seg = LineString([a, b])
            p = seg.interpolate(seg.project(Point(tx, ty)))
            d = (p.x - tx) ** 2 + (p.y - ty) ** 2
            if d < bd:
                bd, best = d, (p.x, p.y)
    return best if best else (tx, ty)


# ---------------------------------------------------------------- 3ds
def write_3ds(blocks, meta, path: Path) -> None:
    """照 3DMine 选择集格式：每个矿块一条字符串，表头末字段为品位类型名。"""
    lines = ["%s, 3DMine String File" % path, "file_version=3DMine_2009"]
    sid = 0
    for b in blocks:
        if b.polygon is None:
            continue
        sid += 1
        lines.append("%d,7,0,1.00,1.00,1.00,0,0.5,0,Continuous,%s" % (sid, b.label))
        lines.append("0,0,0,0,0,0,0,0,")
        polys = list(b.polygon.geoms) if b.polygon.geom_type == "MultiPolygon" else [b.polygon]
        for p in polys:
            for x, y in p.exterior.coords:
                lines.append("1,%.6f,%.6f,%.6f," % (y, x, 0.0))   # 3ds 顺序 N,E,Z
            lines.append("0,0.000,0.000,0.000,0,,1.00,0.000,0.00,1.00,1.00")
    path.write_text("\n".join(lines), encoding="gbk")
