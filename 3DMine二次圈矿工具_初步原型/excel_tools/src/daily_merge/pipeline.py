"""单日数据并入总表 —— 编排层。

一次调用完成四件事：

1. 解析计划（自动探测单日 / 总表文件，支持显式覆盖）；
2. 备份会被改动的总表（``*_并入<日>.bak``，不覆盖旧备份）；
3. 依次并入 xls / mdb / 报告；
4. **顺手把二次圈矿矿块图复制进总目录的「6、…二次圈矿矿块图」文件夹**（去掉「5、」前缀），
   最后复核行数与判重结果。
"""

from __future__ import annotations

import os
import shutil
from dataclasses import replace
from typing import Any

from . import db_mdb, db_xls, map_copy, report_xlsx
from .discover import (DailyMergeOptions, DailyMergePlan, find_map_folder,
                       find_master_files, resolve_plan)
from .schema import (DB_KEY_COLUMNS, DB_TABLES, KIND_DWG, KIND_LABELS,
                     KIND_MDB, KIND_REPORT, KIND_XLS, MERGE_KINDS,
                     map_target_name, normalize_key)

__all__ = ["DailyMergeOptions", "DailyMergePlan", "resolve_plan", "run",
           "backup_file", "verify", "summarize", "restore"]


def backup_file(path: str, day_code: str = "") -> str:
    """把总表备份成 ``<原名>_并入<日>.bak``（已存在则加序号，绝不覆盖）。"""
    stem, ext = os.path.splitext(path)
    tag = "_并入%s" % day_code if day_code else "_备份"
    candidate = "%s%s.bak" % (stem, tag)
    index = 2
    while os.path.exists(candidate):
        candidate = "%s%s-%d.bak" % (stem, tag, index)
        index += 1
    shutil.copy2(path, candidate)
    return candidate


def restore(master_dir: str, day_code: str = None) -> dict:
    """把三份总表还原到 ``*_并入<日>.bak`` 备份。

    ``day_code`` 留空则取该总表最近的一个备份。
    """
    import glob
    import time

    if not master_dir or not os.path.isdir(master_dir):
        raise FileNotFoundError("总表目录不存在：%s" % master_dir)

    restored: dict[str, str] = {}
    for kind, master in find_master_files(master_dir).items():
        if not master or not os.path.isfile(master):
            continue
        stem, ext = os.path.splitext(master)
        if day_code:
            pattern = "%s_并入%s*.bak" % (stem, day_code)
        else:
            pattern = "%s_并入*.bak" % stem
        candidates = sorted(glob.glob(pattern), key=os.path.getmtime)
        if not candidates:
            continue
        source = candidates[-1]                       # 最近的备份
        keep = "%s_还原前%s.bak" % (stem, time.strftime("%Y%m%d-%H%M%S"))
        shutil.copy2(master, keep)
        shutil.copy2(source, master)
        restored[kind] = "%s  <-  %s" % (os.path.basename(master), os.path.basename(source))

    # 矿块图：只在明确给了日期代号时才还原 —— 6、文件夹里还躺着别的日期的 .bak，别误动
    if day_code:
        map_dir = find_map_folder(master_dir)
        if map_dir:
            for source in sorted(glob.glob(os.path.join(map_dir, "*%s.bak" % day_code)),
                                 key=os.path.getmtime):
                stem = source[: -len(".bak")]
                target = stem + ".dwg"
                if not os.path.isfile(target):
                    continue
                shutil.copy2(target, "%s_还原前%s.bak"
                             % (stem, time.strftime("%Y%m%d-%H%M%S")))
                shutil.copy2(source, target)
                restored[KIND_DWG] = "%s  <-  %s" % (os.path.basename(target),
                                                     os.path.basename(source))
    return restored


def _merge_one(kind: str, src: str, dst: str, options: DailyMergeOptions) -> dict:
    if kind == KIND_XLS:
        return db_xls.merge_xls(src, dst,
                                skip_duplicates=options.skip_duplicates,
                                dry_run=options.dry_run)
    if kind == KIND_MDB:
        return db_mdb.merge_mdb(src, dst,
                                skip_duplicates=options.skip_duplicates,
                                dry_run=options.dry_run)
    if kind == KIND_REPORT:
        return report_xlsx.merge_report(src, dst,
                                        with_appendix=options.with_report_appendix,
                                        skip_duplicates=options.skip_duplicates,
                                        dry_run=options.dry_run)
    if kind == KIND_DWG:
        return map_copy.copy_map(src, dst, backup=options.backup, dry_run=options.dry_run)
    raise ValueError("未知的数据类型：%s" % kind)


def run(options: DailyMergeOptions) -> dict:
    """执行并入，返回完整的过程与结果信息。"""
    plan = resolve_plan(options)
    if not plan.executable:
        raise FileNotFoundError(
            "没有可执行的文件对，请检查单日目录 / 总表目录：\n  " + "\n  ".join(plan.missing))

    report: dict[str, Any] = {
        "day_code": plan.day_code,
        "dry_run": options.dry_run,
        "plan": plan,
        "backups": {},
        "results": {},
        "errors": {},
        "warnings": list(plan.warnings),
    }

    # 0) 先查"源文件里到底装的是哪一天"：复制文件夹做模板很常见，
    #    xls/mdb 忘了重新导出 → 里面的孔是前一天的，判重会全跳过（新增 0 行）。
    blocked = _date_mismatch(plan)
    for kind, (dates, total) in blocked.items():
        detail = "、".join("%s×%d" % (d, n) for d, n in dates.most_common())
        report["warnings"].append(
            "单日 %s 里没有任何「%s」的孔（实际是 %s，共 %d 个），与目录日期不一致 —— "
            "很可能是复制文件夹当模板后没有重新导出。这一栏**已跳过**，请重新导出后再跑。"
            % (KIND_LABELS[kind], plan.day_code, detail, total))
    pairs = [p for p in plan.pairs if p[0] not in blocked]

    # 1) 先全面预演一遍，确认哪些总表真的需要改动
    preview_options = replace(options, dry_run=True)
    preview: dict[str, dict] = {}
    for kind, src, dst in pairs:
        try:
            preview[kind] = _merge_one(kind, src, dst, preview_options)
        except Exception as exc:  # noqa: BLE001 - 单类失败不阻断其余
            report["errors"][kind] = "%s" % exc
    report["preview"] = preview

    # 单日报告的"合计 = 各明细之和"自检（只报告，不改源数据）
    if plan.src.get(KIND_REPORT):
        try:
            report["totals_check"] = report_xlsx.check_totals(plan.src[KIND_REPORT])
        except Exception as exc:  # noqa: BLE001
            report["totals_check"] = {"checked": False, "reason": "%s" % exc}

    if options.dry_run:
        report["results"] = dict(preview)
        report["results"].update(_blocked_results(plan, blocked))
        return report

    pending = [p for p in pairs if preview.get(p[0], {}).get("added")]

    # 2) 备份"确实会被改动"的总表（无改动就完全不落备份文件）
    #    矿块图的目标是文件夹，而且复制时自带 <原名>.bak 备份，这里不参与
    if options.backup:
        for kind, _src, dst in pending:
            if kind not in MERGE_KINDS:
                continue
            report["backups"][kind] = backup_file(dst, plan.day_code)

    # 3) 真正写盘；没有改动的直接沿用预演结果
    changed = {p[0] for p in pending}
    for kind, src, dst in pairs:
        if kind in report["errors"]:
            continue
        if kind not in changed:
            report["results"][kind] = preview[kind]
            continue
        try:
            report["results"][kind] = _merge_one(kind, src, dst, options)
        except Exception as exc:  # noqa: BLE001
            report["errors"][kind] = "%s" % exc

    # 被日期检查拦下的那几类，明确给一条"跳过"结果（而不是空着让人以为并过了）
    report["results"].update(_blocked_results(plan, blocked))

    # 4) 收集各类并入返回的提示（不阻断流程，但要让用户看见）
    for res in report["results"].values():
        for notice in res.get("notices", []) or []:
            if notice not in report["warnings"]:
                report["warnings"].append(notice)

    report["verify"] = verify(plan)
    return report


def _blocked_results(plan: DailyMergePlan, blocked: dict[str, tuple]) -> dict[str, dict]:
    """把"日期对不上"的类别表示成一条 skipped 结果，供汇总/界面显示。"""
    out: dict[str, dict] = {}
    for kind, (dates, _total) in blocked.items():
        out[kind] = {
            "kind": kind,
            "src": plan.src.get(kind),
            "dst": plan.dst.get(kind),
            "written": False,
            "skipped": True,
            "added": 0,
            "reason": "单日文件里的孔不是「%s」的（实际 %s），已跳过"
                      % (plan.day_code,
                         "、".join("%s×%d" % (d, n) for d, n in dates.most_common())),
        }
    return out


def _source_dates(kind: str, path: str):
    """统计单日文件里孔的日期段（取工程号第三段），返回 Counter。"""
    import re
    from collections import Counter

    if kind == KIND_XLS:
        info = db_xls.read_sheets(path).get("定位表")
        holes = [str(r.get("工程号")) for r in info["rows"]] if info else []
    elif kind == KIND_MDB:
        conn = db_mdb.connect(path, write=False)
        try:
            holes = [str(r.get("工程号"))
                     for r in db_mdb.read_table_rows(conn, "定位表", ["工程号"])]
        finally:
            conn.close()
    else:
        return Counter()

    pattern = re.compile(r"[A-Za-z]{2,3}-\d+-(\d{3,4})-")
    dates = Counter()
    for hole in holes:
        match = pattern.search(hole)
        if match:
            dates[match.group(1)] += 1
    return dates


def _date_mismatch(plan: DailyMergePlan) -> dict[str, tuple]:
    """找"源文件里一个当天的孔都没有"的类别 —— 这种基本是源文件没重新导出。

    只在**完全没有**当天数据时才拦；某天里夹着前几天没爆完的孔（很常见）不拦。
    """
    day_digits = "".join(ch for ch in (plan.day_code or "") if ch.isdigit())
    if not day_digits:
        return {}
    blocked: dict[str, tuple] = {}
    for kind in (KIND_XLS, KIND_MDB):
        path = plan.src.get(kind)
        if not path or not os.path.isfile(path):
            continue
        try:
            dates = _source_dates(kind, path)
        except Exception:  # noqa: BLE001 - 读不出来就别拦，交给后面的并入去报错
            continue
        if not dates or day_digits in dates:
            continue
        blocked[kind] = (dates, sum(dates.values()))
    return blocked


# --------------------------------------------------------------------------
# 复核
# --------------------------------------------------------------------------

def _xls_holes(path: str) -> list[str]:
    book = db_xls.read_xls_sheets(path)
    info = book.get("定位表")
    return [str(r.get("工程号")) for r in info["rows"]] if info else []


def _mdb_holes(path: str) -> list[str]:
    conn = db_mdb.connect(path, write=False)
    try:
        rows = db_mdb.read_table_rows(conn, "定位表", ["工程号"])
        return [str(r.get("工程号")) for r in rows]
    finally:
        conn.close()


def _report_blocks(path: str) -> tuple[int, list[str]]:
    """数总表第一张表里有几个日报告块（xls / xlsx 都行）。"""
    from ..excel_core.filetype import FMT_XLSX, sniff
    from .schema import REPORT_TOTAL_TOKENS, is_report_title

    totals, titles = 0, []
    if sniff(path) == FMT_XLSX:
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=False, read_only=True)
        try:
            ws = wb[wb.sheetnames[0]]
            for row in ws.iter_rows(min_col=1, max_col=1):
                value = row[0].value
                if value is None:
                    continue
                text = str(value).strip()
                if text in REPORT_TOTAL_TOKENS:
                    totals += 1
                elif is_report_title(text):
                    titles.append(text)
        finally:
            wb.close()
    else:
        import xlrd

        book = xlrd.open_workbook(path)
        try:
            sh = book.sheet_by_index(0)
            for r in range(sh.nrows):
                value = sh.cell_value(r, 0)
                if value is None or str(value).strip() == "":
                    continue
                text = str(value).strip()
                if text in REPORT_TOTAL_TOKENS:
                    totals += 1
                elif is_report_title(text):
                    titles.append(text)
        finally:
            book.release_resources()
    return totals, titles


def verify(plan: DailyMergePlan) -> dict:
    """复核并入结果：孔数、日期分段、重复情况、报告块数量。"""
    out: dict[str, Any] = {}

    for kind, reader in ((KIND_XLS, _xls_holes), (KIND_MDB, _mdb_holes)):
        path = plan.dst.get(kind)
        if not path or not os.path.isfile(path):
            continue
        try:
            holes = reader(path)
        except Exception as exc:  # noqa: BLE001
            out[kind] = {"error": str(exc)}
            continue
        out[kind] = _hole_summary(holes)

    path = plan.dst.get(KIND_REPORT)
    if path and os.path.isfile(path):
        try:
            totals, titles = _report_blocks(path)
            out[KIND_REPORT] = {"total_rows": totals, "blocks": len(titles), "titles": titles}
        except Exception as exc:  # noqa: BLE001
            out[KIND_REPORT] = {"error": str(exc)}

    src_map, map_dir = plan.src.get(KIND_DWG), plan.dst.get(KIND_DWG)
    if src_map and map_dir:
        target = os.path.join(map_dir, map_target_name(src_map))
        exists = os.path.isfile(target)
        out[KIND_DWG] = {
            "target": target,
            "exists": exists,
            "identical": map_copy.files_identical(src_map, target) if exists else False,
        }

    return out


def _hole_summary(holes: list[str]) -> dict:
    import re
    from collections import Counter

    segments = Counter()
    for hole in holes:
        # 工程号形如 BS-3940-1008-P12 / NS-3970-0719-P1，前缀按"两三个字母-平台号-日期"取
        match = re.match(r"([A-Za-z]{2,3}-\d+-\d+)", str(hole))
        if match:
            segments[match.group(1)] += 1

    seen, dup = set(), []
    for hole in holes:
        key = normalize_key(hole)
        if key in seen:
            dup.append(hole)
        seen.add(key)

    return {
        "total": len(holes),
        "unique": len(seen),
        "duplicates": dup,
        "segments": dict(segments),
        "first": holes[0] if holes else "",
        "last": holes[-1] if holes else "",
    }


def summarize(report: dict) -> list[str]:
    """把 run() 的结果整理成可直接打印的中文行。"""
    lines: list[str] = []
    if report.get("dry_run"):
        lines.append("[预演] 只计算不写盘")
    lines.append("日期代号：%s" % (report.get("day_code") or "(未知)"))

    for kind, info in report["backups"].items():
        lines.append("[备份] %s" % info)

    for kind, res in report["results"].items():
        label = KIND_LABELS[kind]
        if kind == KIND_DWG:
            lines.append(map_copy.describe(res))
            continue
        if res.get("skipped") and "tables" not in res:
            lines.append("[跳过] %s：%s" % (label, res.get("reason", "已跳过")))
            continue
        if kind == KIND_REPORT:
            if res.get("skipped"):
                lines.append("[跳过] %s：%s" % (label, res["reason"]))
            else:
                lines.append("[并入] %s：追加 %d 行（标题「%s」，公式 %d 个，起始第 %d 行）"
                             % (label, res["block_rows"], res["title"], res["formulas"], res["insert_at"]))
            continue
        pieces = []
        for table, st in res["tables"].items():
            if st["status"] != "ok":
                pieces.append("%s 跳过(%s)" % (table, st["reason"]))
            else:
                pieces.append("%s %d→%d(+%d)" % (table, st["existing"], st["existing"] + st["added"], st["added"]))
        lines.append("[并入] %s：新增 %d 行；%s" % (label, res["added"], "，".join(pieces)))

    for kind, err in report["errors"].items():
        lines.append("[失败] %s：%s" % (KIND_LABELS[kind], err))

    verify_info = report.get("verify") or {}
    for kind in (KIND_XLS, KIND_MDB):
        info = verify_info.get(kind)
        if not info or "error" in info:
            continue
        seg = "、".join("%s×%d" % (k, v) for k, v in sorted(info["segments"].items()))
        dup = ("，重复 %d 个" % len(info["duplicates"])) if info["duplicates"] else ""
        lines.append("[复核] %s：共 %d 孔（%s）%s"
                     % (KIND_LABELS[kind], info["total"], seg, dup))
    info = verify_info.get(KIND_REPORT)
    if info and "error" not in info:
        lines.append("[复核] %s：%d 个日报告块，%d 个合计行"
                     % (KIND_LABELS[KIND_REPORT], info["blocks"], info["total_rows"]))

    info = verify_info.get(KIND_DWG)
    if info:
        if not info.get("exists"):
            lines.append("[复核] %s：6、文件夹里还没有 %s"
                         % (KIND_LABELS[KIND_DWG], os.path.basename(info["target"])))
        elif info.get("identical"):
            lines.append("[复核] %s：6、文件夹里的 %s 与单日图一致 ✓"
                         % (KIND_LABELS[KIND_DWG], os.path.basename(info["target"])))
        else:
            lines.append("[复核] %s：6、文件夹里的 %s 与单日图内容不同！"
                         % (KIND_LABELS[KIND_DWG], os.path.basename(info["target"])))

    check = report.get("totals_check") or {}
    if check.get("checked"):
        if check["mismatches"]:
            for item in check["mismatches"]:
                lines.append("[警告] 单日报告的「%s」合计与各明细之和差 %s（明细 %s / 合计 %s）"
                             "—— 源报告本身公式疑有误，已按原样并入，请人工核对。"
                             % (item["column"], item["diff"], item["detail_sum"], item["total"]))
        else:
            lines.append("[校验] 单日报告合计与各明细之和一致 ✓")
    elif check:
        lines.append("[提示] 单日报告合计未校验：%s" % check.get("reason", ""))

    for warn in report.get("warnings", []):
        lines.append("[提示] %s" % warn)
    return lines
