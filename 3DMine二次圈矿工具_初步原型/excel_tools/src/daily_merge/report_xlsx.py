"""炮孔数据报告（总表）的并入 —— 目标是 ``.xlsx`` 或 ``.xls`` 都支持。

总表是"从上到下依次堆叠的日报告块"：
每一块 = 标题行 + 表头行 + 各品位档明细 + 小计 + 合计。

并入时取单日报告的**汇总块**（标题行 → 合计行），在总表末尾空一行后整块粘贴，
块内的公式按行偏移整体平移（``=F5/D5*100`` → ``=F27/D27*100``）。

格式按**文件内容**判断（不看扩展名），三种组合都走通：

============================  ==========================  ==============================
总表                          单日报告                    做法
============================  ==========================  ==============================
``.xlsx``（北部山头）          ``.xlsx``                   openpyxl：值 + 公式 + 字体边框合并单元格
``.xls``（南部山头）           ``.xlsx``                   xlutils：值 + 公式，样式借总表同类行
``.xls``                      ``.xls``                    同上，只是源里的公式已只剩计算值
============================  ==========================  ==============================
"""

from __future__ import annotations

import copy as _copy
import os
import re
from dataclasses import dataclass, field
from typing import Any

import xlrd
import xlwt
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from xlutils.filter import XLRDReader, XLWTWriter, process

from ..excel_core.filetype import FMT_XLS, FMT_XLSX, sniff
from .schema import (KINDS, REPORT_SUBTOTAL_TOKENS, REPORT_TOTAL_TOKEN,
                     REPORT_TOTAL_TOKENS, REPORT_TYPE_COLUMN, is_report_title)

__all__ = ["find_summary_block", "existing_titles", "check_totals", "merge_report"]

_BLANK_SEPARATOR = 1


# --------------------------------------------------------------------------
# 小工具
# --------------------------------------------------------------------------

def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _cell_value(sheet, row: int, col: int) -> Any:
    """xlrd 单元格 -> Python 值（空 -> None）。"""
    cell = sheet.cell(row, col)
    if cell.ctype in (0, 6):          # empty / blank
        return None
    if cell.ctype == 1:               # text
        return cell.value if cell.value != "" else None
    if cell.ctype in (2, 3):          # number / date
        return cell.value
    if cell.ctype == 4:               # boolean
        return bool(cell.value)
    if cell.ctype == 5:               # error
        return None
    return None


def _last_content_row(ws) -> int:
    last = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and str(cell.value).strip() != "":
                last = max(last, cell.row)
    return last


def _block_columns(ws, r1: int, r2: int) -> int:
    """块内实际用到的最大列号。"""
    ncol = 1
    for row in ws.iter_rows(min_row=r1, max_row=r2):
        for cell in row:
            if cell.value is not None and str(cell.value).strip() != "":
                ncol = max(ncol, cell.column)
    return ncol


def find_summary_block(ws) -> tuple[int, int]:
    """定位汇总块的首尾行（含标题行与合计行）。找不到抛 ValueError。"""
    first = None
    total = None
    for row in ws.iter_rows(min_col=1, max_col=1):
        text = _text(row[0].value)
        if first is None and text:
            first = row[0].row
        if text in REPORT_TOTAL_TOKENS:
            total = row[0].row
            break
    if total is None:
        raise ValueError("单日报告里没找到「%s」行，无法确定汇总块范围"
                         % "/".join(REPORT_TOTAL_TOKENS))
    if first is None:
        first = 1
    return first, total


def existing_titles(ws) -> list[str]:
    """总表里已有的日报告块标题（用于防止同一日重复并入）。"""
    titles = []
    for row in ws.iter_rows(min_col=1, max_col=1):
        value = row[0].value
        if is_report_title(value):
            titles.append(_text(value))
    return titles


# --------------------------------------------------------------------------
# 汇总块：统一读成"行值 + 合并单元格"
# --------------------------------------------------------------------------

@dataclass
class Block:
    """一个单日报告的汇总块（源里是第 r1~r2 行，0-based）。"""

    title: str = ""
    values: list[list[Any]] = field(default_factory=list)   # [行][列]
    merged: list[tuple[int, int, int, int]] = field(default_factory=list)  # (r0,c0,r1,c1) 0-based 含
    ncol: int = 1
    r1: int = 0
    r2: int = 0
    has_formula: bool = False

    @property
    def height(self) -> int:
        return len(self.values)


def _read_block(src_path: str, *, with_appendix: bool = False) -> Block:
    """按真实格式读源报告的汇总块。"""
    if sniff(src_path) == FMT_XLSX:
        return _read_block_xlsx(src_path, with_appendix=with_appendix)
    return _read_block_xls(src_path, with_appendix=with_appendix)


def _read_block_xlsx(src_path: str, *, with_appendix: bool = False) -> Block:
    wb = load_workbook(src_path, data_only=False)
    try:
        ws = wb[wb.sheetnames[0]]
        r1, r2 = find_summary_block(ws)
        if with_appendix:
            r2 = max(r2, _last_content_row(ws))
        ncol = _block_columns(ws, r1, r2)

        values = [[ws.cell(r, c).value for c in range(1, ncol + 1)]
                  for r in range(r1, r2 + 1)]
        merged = []
        for rng in ws.merged_cells.ranges:
            if rng.min_row >= r1 and rng.max_row <= r2 and rng.max_col <= ncol:
                merged.append((rng.min_row - 1, rng.min_col - 1,
                               rng.max_row - 1, rng.max_col - 1))
        has_formula = any(isinstance(v, str) and v.startswith("=")
                          for row in values for v in row)
        return Block(title=_text(ws.cell(r1, 1).value), values=values, merged=merged,
                     ncol=ncol, r1=r1 - 1, r2=r2 - 1, has_formula=has_formula)
    finally:
        wb.close()


def _read_block_xls(src_path: str, *, with_appendix: bool = False) -> Block:
    book = xlrd.open_workbook(src_path)
    try:
        sh = book.sheet_by_index(0)
        first = total = None
        for r in range(sh.nrows):
            text = _text(sh.cell_value(r, 0))
            if first is None and text:
                first = r
            if text in REPORT_TOTAL_TOKENS:
                total = r
                break
        if total is None:
            raise ValueError("单日报告里没找到「%s」行，无法确定汇总块范围"
                             % "/".join(REPORT_TOTAL_TOKENS))
        if first is None:
            first = 0
        last = max(total, sh.nrows - 1) if with_appendix else total

        ncol = 1
        for r in range(first, last + 1):
            for c in range(sh.ncols):
                if _text(sh.cell_value(r, c)):
                    ncol = max(ncol, c + 1)

        values = [[_cell_value(sh, r, c) for c in range(ncol)]
                  for r in range(first, last + 1)]
        merged = []
        for rlo, rhi, clo, chi in sh.merged_cells:
            if rlo >= first and rhi - 1 <= last and chi - 1 <= ncol - 1:
                merged.append((rlo, clo, rhi - 1, chi - 1))
        return Block(title=_text(sh.cell_value(first, 0)), values=values, merged=merged,
                     ncol=ncol, r1=first, r2=last, has_formula=False)
    finally:
        book.release_resources()


# --------------------------------------------------------------------------
# 合计自检（两种格式都能查）
# --------------------------------------------------------------------------

def _read_grid(path: str) -> list[list[Any]]:
    """把第一张表读成二维值数组（不关心格式）。"""
    if sniff(path) == FMT_XLSX:
        wb = load_workbook(path, data_only=True)
        try:
            ws = wb[wb.sheetnames[0]]
            return [list(row) for row in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
    book = xlrd.open_workbook(path)
    try:
        sh = book.sheet_by_index(0)
        return [[_cell_value(sh, r, c) for c in range(sh.ncols)]
                for r in range(sh.nrows)]
    finally:
        book.release_resources()


# 合计行所在列（0-based）：体积 / 重量 / 金属量（平均品位是推导列，不参与求和校验）
_SUM_COLUMNS = ((2, "体积"), (3, "重量"), (5, "金属量"))


def check_totals(src_path: str, tolerance: float = 0.05) -> dict:
    """用文件里缓存的计算值，复核"合计 = 各明细行之和"。

    只做比对并报告差异，绝不改动源数据 —— 源报告本身写错公式时，
    这里会明确提示，避免错误被原样搬进总表。
    """
    grid = _read_grid(src_path)
    if not grid:
        return {"checked": False, "reason": "表是空的"}

    first = total = None
    for r, row in enumerate(grid):
        text = _text(row[0] if row else None)
        if first is None and text:
            first = r
        if text in REPORT_TOTAL_TOKENS:
            total = r
            break
    if total is None:
        return {"checked": False,
                "reason": "单日报告里没找到「%s」行，无法确定汇总块范围"
                          % "/".join(REPORT_TOTAL_TOKENS)}
    if first is None:
        first = 0

    sums = {name: 0.0 for _col, name in _SUM_COLUMNS}
    details = 0
    for r in range(first + 1, total):
        row = grid[r]
        label = _text(row[0] if row else None)
        if not label or label in REPORT_TOTAL_TOKENS or label in REPORT_SUBTOTAL_TOKENS:
            continue
        details += 1
        for col, name in _SUM_COLUMNS:
            value = row[col] if col < len(row) else None
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                sums[name] += float(value)

    total_row = grid[total]
    totals = {}
    for col, name in _SUM_COLUMNS:
        value = total_row[col] if col < len(total_row) else None
        totals[name] = (float(value)
                        if isinstance(value, (int, float)) and not isinstance(value, bool)
                        else None)

    mismatches = []
    for _col, name in _SUM_COLUMNS:
        got = totals[name]
        if got is None:
            continue
        if abs(got - sums[name]) > tolerance:
            mismatches.append({"column": name, "detail_sum": round(sums[name], 3),
                               "total": round(got, 3),
                               "diff": round(got - sums[name], 3)})
    return {
        "checked": True,
        "detail_rows": details,
        "total_label": _text(total_row[0] if total_row else None),
        "detail_sum": {k: round(v, 3) for k, v in sums.items()},
        "total": {k: (round(v, 3) if v is not None else None) for k, v in totals.items()},
        "mismatches": mismatches,
    }


# --------------------------------------------------------------------------
# 写入：xlsx 总表（原有的整块粘贴）
# --------------------------------------------------------------------------

def _copy_style(src_cell, dst_cell) -> None:
    dst_cell.font = _copy.copy(src_cell.font)
    dst_cell.border = _copy.copy(src_cell.border)
    dst_cell.fill = _copy.copy(src_cell.fill)
    dst_cell.alignment = _copy.copy(src_cell.alignment)
    dst_cell.protection = _copy.copy(src_cell.protection)
    dst_cell.number_format = src_cell.number_format


def _translate(value: Any, origin: str, target: str) -> Any:
    if isinstance(value, str) and value.startswith("="):
        try:
            return Translator(value, origin=origin).translate_formula(target)
        except Exception:  # noqa: BLE001 - 平移失败就原样保留
            return value
    return value


def _merge_into_xlsx(src_path: str, dst_path: str, *,
                     with_appendix: bool = False,
                     separator: int = _BLANK_SEPARATOR,
                     skip_duplicates: bool = True,
                     dry_run: bool = False) -> dict:
    src_wb = load_workbook(src_path, data_only=False)
    try:
        src_ws = src_wb[src_wb.sheetnames[0]]

        r1, r2 = find_summary_block(src_ws)
        if with_appendix:
            r2 = max(r2, _last_content_row(src_ws))
        ncol = _block_columns(src_ws, r1, r2)
        title = _text(src_ws.cell(r1, 1).value)

        dst_wb = load_workbook(dst_path, data_only=False)
        try:
            dst_ws = dst_wb[dst_wb.sheetnames[0]]
            dst_last = _last_content_row(dst_ws)

            if skip_duplicates and title and title in existing_titles(dst_ws):
                return _skipped_result(src_path, dst_path, title, r2 - r1 + 1,
                                       ncol, dst_last)

            start = dst_last + 1 + max(separator, 0)
            delta = start - r1

            formulas = sum(
                1
                for row in src_ws.iter_rows(min_row=r1, max_row=r2, max_col=ncol)
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")
            )

            if not dry_run:
                for row in src_ws.iter_rows(min_row=r1, max_row=r2, max_col=ncol):
                    for src_cell in row:
                        if src_cell.value is None:
                            continue
                        dst_cell = dst_ws.cell(src_cell.row + delta, src_cell.column)
                        dst_cell.value = _translate(src_cell.value, src_cell.coordinate,
                                                    dst_cell.coordinate)
                        _copy_style(src_cell, dst_cell)
                # 合并单元格（只搬块内的）
                for rng in src_ws.merged_cells.ranges:
                    if rng.min_row >= r1 and rng.max_row <= r2 and rng.max_col <= ncol:
                        dst_ws.merge_cells(
                            start_row=rng.min_row + delta, start_column=rng.min_col,
                            end_row=rng.max_row + delta, end_column=rng.max_col,
                        )
                # 行高
                for r in range(r1, r2 + 1):
                    height = src_ws.row_dimensions[r].height
                    if height:
                        dst_ws.row_dimensions[r + delta].height = height

                try:      # 让 Excel 打开时重算（openpyxl 保存会清掉公式的缓存值）
                    dst_wb.calculation.fullCalcOnLoad = True
                except Exception:  # noqa: BLE001
                    pass
                dst_wb.save(dst_path)
        finally:
            dst_wb.close()
    finally:
        src_wb.close()

    return {
        "kind": KINDS[2],
        "src": os.path.abspath(src_path),
        "dst": os.path.abspath(dst_path),
        "written": not dry_run,
        "skipped": False,
        "title": title,
        "block_rows": r2 - r1 + 1,
        "columns": ncol,
        "formulas": formulas,
        "insert_at": start,
        "before_rows": dst_last,
        "after_rows": dst_last + (0 if dry_run else (separator + r2 - r1 + 1)),
        "added": r2 - r1 + 1,
        "format": {"src": sniff(src_path), "dst": "xlsx"},
    }


def _skipped_result(src_path: str, dst_path: str, title: str, block_rows: int,
                    ncol: int, dst_last: int) -> dict:
    return {
        "kind": KINDS[2],
        "src": os.path.abspath(src_path),
        "dst": os.path.abspath(dst_path),
        "written": False,
        "skipped": True,
        "title": title,
        "block_rows": block_rows,
        "columns": ncol,
        "formulas": 0,
        "insert_at": 0,
        "before_rows": dst_last,
        "after_rows": dst_last,
        "added": 0,
        "reason": "总表中已存在「%s」，未重复并入" % title,
        "format": {"src": sniff(src_path), "dst": sniff(dst_path)},
    }


# --------------------------------------------------------------------------
# 写入：xls（BIFF）总表
# --------------------------------------------------------------------------

_REF_RE = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")


def _shift_formula(formula: str, delta: int) -> str:
    """把公式里的行号整体下移 ``delta`` 行（绝对行 ``$5`` 不动）。"""
    if not delta:
        return formula

    def repl(match: re.Match) -> str:
        col_abs, col, row_abs, row = match.groups()
        if row_abs:
            return match.group(0)
        return "%s%s%s%d" % (col_abs, col, row_abs, int(row) + delta)

    return _REF_RE.sub(repl, formula)


def _role_of(value_txt: str) -> str:
    """判断一行在报告块里扮演什么角色，用来借总表里同类行的样式。"""
    text = _text(value_txt)
    if not text:
        return "blank"
    if REPORT_TYPE_COLUMN in text:
        return "header"
    if text in REPORT_SUBTOTAL_TOKENS:
        return "subtotal"
    if text in REPORT_TOTAL_TOKENS:
        return "total"
    if is_report_title(text):
        return "title"
    return "data"


def _role_rows(sheet, last_row: int) -> dict[str, int]:
    """总表里每种角色最后一次出现的行号（0-based）。"""
    rows: dict[str, int] = {}
    for r in range(0, last_row + 1):
        role = _role_of(sheet.cell_value(r, 0))
        if role != "blank":
            rows[role] = r
    return rows


def _merge_into_xls(src_path: str, dst_path: str, *,
                    with_appendix: bool = False,
                    separator: int = _BLANK_SEPARATOR,
                    skip_duplicates: bool = True,
                    dry_run: bool = False) -> dict:
    block = _read_block(src_path, with_appendix=with_appendix)

    book = xlrd.open_workbook(dst_path, formatting_info=True)
    try:
        dst_sheet = book.sheet_by_index(0)
        dst_last = -1
        for r in range(dst_sheet.nrows - 1, -1, -1):
            if any(_text(dst_sheet.cell_value(r, c)) for c in range(dst_sheet.ncols)):
                dst_last = r
                break

        titles = [_text(dst_sheet.cell_value(r, 0))
                  for r in range(0, dst_last + 1)
                  if is_report_title(_text(dst_sheet.cell_value(r, 0)))]
        if skip_duplicates and block.title and block.title in titles:
            return _skipped_result(src_path, dst_path, block.title, block.height,
                                   block.ncol, dst_last + 1)

        start = dst_last + 1 + max(separator, 0)
        delta = start - block.r1          # 源块首行 -> 目标首行 的行偏移

        formulas = sum(1 for row in block.values for v in row
                       if isinstance(v, str) and v.startswith("="))

        if not dry_run:
            writer = XLWTWriter()
            process(XLRDReader(book, os.path.basename(dst_path)), writer)
            wt_book = writer.output[0][1]
            wt_sheet = wt_book.get_sheet(0)

            role_rows = _role_rows(dst_sheet, dst_last)
            fallback = role_rows.get("data", max(dst_last, 0))

            def style_at(style_row: int, col: int):
                c = min(col, max(dst_sheet.ncols - 1, 0))
                xf = dst_sheet.cell(max(style_row, 0), c).xf_index
                return (writer.style_list[xf]
                        if xf < len(writer.style_list) else None)

            def is_merged_child(row_i: int, col_i: int) -> bool:
                for m_r1, m_c1, m_r2, m_c2 in block.merged:
                    if (m_r1 <= row_i <= m_r2 and m_c1 <= col_i <= m_c2
                            and (row_i, col_i) != (m_r1, m_c1)):
                        return True
                return False

            for i, row in enumerate(block.values):
                r = start + i
                role = _role_of(row[0] if row else None)
                style_row = role_rows.get(role, fallback)
                for c, value in enumerate(row):
                    if value is None or (isinstance(value, str) and value == ""):
                        continue
                    if is_merged_child(i, c):
                        continue
                    if isinstance(value, str) and value.startswith("="):
                        try:
                            value = xlwt.Formula(_shift_formula(value[1:], delta))
                        except Exception:  # noqa: BLE001
                            pass
                    style = style_at(style_row, c)
                    if style is not None:
                        wt_sheet.write(r, c, value, style)
                    else:
                        wt_sheet.write(r, c, value)

            # 合并单元格：块内标题行常常整行合并
            for m_r1, m_c1, m_r2, m_c2 in block.merged:
                r1 = start + (m_r1 - block.r1)
                r2 = start + (m_r2 - block.r1)
                value = block.values[m_r1 - block.r1][m_c1]
                style = style_at(role_rows.get("title", fallback), m_c1)
                if style is not None:
                    wt_sheet.write_merge(r1, r2, m_c1, m_c2,
                                         "" if value is None else value, style)
                else:
                    wt_sheet.write_merge(r1, r2, m_c1, m_c2,
                                         "" if value is None else value)

            wt_book.save(dst_path)
    finally:
        book.release_resources()

    return {
        "kind": KINDS[2],
        "src": os.path.abspath(src_path),
        "dst": os.path.abspath(dst_path),
        "written": not dry_run,
        "skipped": False,
        "title": block.title,
        "block_rows": block.height,
        "columns": block.ncol,
        "formulas": formulas,
        "insert_at": start + 1,
        "before_rows": dst_last + 1,
        "after_rows": dst_last + 1 + (0 if dry_run else (separator + block.height)),
        "added": block.height,
        "format": {"src": sniff(src_path), "dst": "xls"},
    }


# --------------------------------------------------------------------------
# 对外入口
# --------------------------------------------------------------------------

def merge_report(src_path: str, dst_path: str, *,
                 with_appendix: bool = False,
                 separator: int = _BLANK_SEPARATOR,
                 skip_duplicates: bool = True,
                 dry_run: bool = False) -> dict:
    """把 ``src_path`` 单日报告的汇总块追加到 ``dst_path`` 总表末尾。

    总表是 ``.xlsx`` 还是 ``.xls`` 由**文件内容**判断，两种都支持。
    """
    kind = sniff(dst_path)
    if kind == FMT_XLS:
        return _merge_into_xls(src_path, dst_path, with_appendix=with_appendix,
                               separator=separator, skip_duplicates=skip_duplicates,
                               dry_run=dry_run)
    if kind == FMT_XLSX:
        return _merge_into_xlsx(src_path, dst_path, with_appendix=with_appendix,
                                separator=separator, skip_duplicates=skip_duplicates,
                                dry_run=dry_run)
    raise ValueError("总表不是 Excel 文件（扩展名 %s）：%s"
                     % (os.path.splitext(dst_path)[1], dst_path))
