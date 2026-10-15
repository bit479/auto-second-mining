"""自检：单日数据并入总表。

流程（全程在 ``_out/daily_test`` 的副本里操作，绝不碰真实数据）：

1. 用 3940 平台 1008 的三个单日文件 + 三份总表搭一个沙箱；
2. 跑一次并入，逐行比对 xls / mdb 的追加行、检查报告块的公式平移；
3. 再跑一次，确认"判重"生效（新增 0 行、报告不重复粘）；
4. 从备份还原，确认总表回到并入前的孔数。

用法::

    python tests/test_daily_merge.py
"""

from __future__ import annotations

import os
import shutil
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from src.daily_merge.pipeline import (DailyMergeOptions, resolve_plan,  # noqa: E402
                                      restore, run, summarize)
from src.daily_merge.schema import KIND_MDB, KIND_REPORT, KIND_XLS  # noqa: E402

REAL_MASTER = r"D:\WK\N CRT\北部山头3940平台2026"
REAL_DAY = os.path.join(REAL_MASTER, r"7、单日数据\北部山头3940平台-1008")
SANDBOX = os.path.join(ROOT, "_out", "daily_test")

MASTER_FILES = [
    "1、北部山头3940平台 炮孔数据库-2026.xls",
    "2、北部山头3940平台炮孔数据库-2026.mdb",
    "3、北部山头3940平台炮孔数据报告-2026综合.xlsx",
]
DAY_FILES = [
    "1、3940平台 炮孔数据库1008.xls",
    "2、3940平台炮孔数据库1008.mdb",
    "3940平台 1008炮孔数据报告.xlsx",
]

_checks: list[tuple[bool, str]] = []


def check(ok: bool, message: str) -> None:
    _checks.append((bool(ok), message))
    print("  %s  %s" % ("OK  " if ok else "FAIL", message))


def build_sandbox() -> str:
    if os.path.exists(SANDBOX):
        shutil.rmtree(SANDBOX)
    day_dir = os.path.join(SANDBOX, "北部山头3940平台2026", "7、单日数据", "北部山头3940平台-1008")
    os.makedirs(day_dir)
    for name in MASTER_FILES:
        shutil.copy2(os.path.join(REAL_MASTER, name),
                     os.path.join(SANDBOX, "北部山头3940平台2026", name))
    for name in DAY_FILES:
        shutil.copy2(os.path.join(REAL_DAY, name), os.path.join(day_dir, name))
    return day_dir


# --------------------------------------------------------------------------
# 逐行比对
# --------------------------------------------------------------------------

def compare_xls(day_dir: str) -> None:
    import xlrd

    src = xlrd.open_workbook(os.path.join(day_dir, DAY_FILES[0]))
    dst = xlrd.open_workbook(os.path.join(SANDBOX, "北部山头3940平台2026", MASTER_FILES[0]))
    check(dst.sheet_names() == ["定位表", "测斜表", "pit_blast"],
          "xls 保留原有 3 张表（不含单日的多余 Sheet1）")
    for table in ("定位表", "测斜表", "pit_blast"):
        a, b = src.sheet_by_name(table), dst.sheet_by_name(table)
        offset = b.nrows - a.nrows
        bad = 0
        for r in range(1, a.nrows):
            for c in range(a.ncols):
                va, vb = a.cell_value(r, c), b.cell_value(r + offset, c)
                if isinstance(va, float):
                    if abs(va - vb) > 1e-9:
                        bad += 1
                elif str(va) != str(vb):
                    bad += 1
        check(bad == 0 and offset == 70,
              "%s：追加 %d 行、逐值一致（表头行 +70）" % (table, a.nrows - 1))


def compare_mdb(day_dir: str) -> None:
    from src.daily_merge.db_mdb import connect, list_tables, read_table_rows

    src = connect(os.path.join(day_dir, DAY_FILES[1]), write=False)
    dst = connect(os.path.join(SANDBOX, "北部山头3940平台2026", MASTER_FILES[1]), write=False)
    try:
        for table, columns in (("定位表", ["工程号", "开孔坐标E", "开孔坐标N", "开孔坐标R", "轨迹类型", "最大孔深"]),
                               ("测斜表", ["工程号", "深度", "方位角", "倾角"]),
                               ("pit_blast", ["工程号", "从", "至", "Au"])):
            a = read_table_rows(src, table, columns)
            b = read_table_rows(dst, table, columns)
            bad = sum(
                1
                for i, row in enumerate(a)
                for col in columns
                if (abs(b[70 + i][col] - row[col]) > 1e-9
                    if isinstance(row[col], float) else str(b[70 + i][col]) != str(row[col]))
            )
            check(len(b) == 120 and bad == 0,
                  "%s：总表 %d 行（70+50），追加行逐值一致" % (table, len(b)))
        check(set(list_tables(dst)) == set(list_tables(src)) | set(list_tables(dst)),
              "mdb 的 3DMine 元数据表未被破坏")
    finally:
        src.close()
        dst.close()


def compare_report() -> None:
    from openpyxl import load_workbook

    from src.daily_merge.report_xlsx import check_totals

    day_dir = os.path.join(SANDBOX, "北部山头3940平台2026", "7、单日数据", "北部山头3940平台-1008")
    ws = load_workbook(os.path.join(SANDBOX, "北部山头3940平台2026", MASTER_FILES[2])).active
    check(ws.cell(23, 1).value == "3940平台 1008炮孔数据报告", "报告块标题落在第 23 行")
    check(ws.cell(24, 1).value == "类型", "第 24 行是表头")
    check(ws.cell(27, 3).value == "=C25+C26", "小计公式按行偏移平移：=C3+C4 -> =C25+C26")
    check(ws.cell(30, 5).value == "=F30/D30*100", "平均品位公式平移：=F8/D8*100 -> =F30/D30*100")
    check(ws.cell(30, 4).value == "=D27+D29", "合计公式平移：=D5+D7 -> =D27+D29")
    check("A23:F23" in {str(r) for r in ws.merged_cells.ranges}, "标题行合并单元格已复制")
    check(ws.cell(25, 3).value == 714.095, "明细数值原样带入")
    check(ws.cell(2, 4).value == "重量（t）", "总表原有内容未被改动")

    totals = check_totals(os.path.join(day_dir, DAY_FILES[2]))
    check(totals.get("checked") and not totals.get("mismatches"),
          "单日报告合计自检：明细之和与合计一致")


# --------------------------------------------------------------------------
# 主流程
# --------------------------------------------------------------------------

def main() -> int:
    print("=" * 78)
    print("沙箱：", SANDBOX)
    day_dir = build_sandbox()

    options = DailyMergeOptions(day_dir=day_dir)
    plan = resolve_plan(options)
    print("-" * 78)
    for line in plan.describe():
        print(line)
    print("-" * 78)

    check(len(plan.pairs) == 3, "自动识别出 3 对文件")

    print("\n[1/4] 首次并入")
    report = run(options)
    for line in summarize(report):
        print("   " + line)
    check(not report["errors"], "没有失败项")
    nx = report["results"][KIND_XLS]["added"]
    nm = report["results"][KIND_MDB]["added"]
    # 真实总表里**可能已经并过这一天**（用户自己并过）——那时"新增 0 行"才是对的。
    # 所以这里两种状态都算通过：真并进去了就查行数，已经并过就查判重确实生效。
    if nx == 0 and nm == 0:
        print("   （总表里这一天已存在——本次验的是判重路径）")
        check(report["results"][KIND_REPORT].get("skipped") or
              report["results"][KIND_REPORT]["added"] == 0, "报告未重复并入")
        check(len(report["backups"]) == 0, "没有改动就不备份")
    else:
        check(nx == 150, "xls 追加 150 行（3 表 × 50）")
        check(nm == 150, "mdb 追加 150 行（3 表 × 50）")
        check(report["results"][KIND_REPORT]["added"] == 8, "报告追加 8 行块")
        check(len(report["backups"]) == 3, "三份总表都已备份")

    print("\n[2/4] 逐行比对")
    compare_xls(day_dir)
    compare_mdb(day_dir)
    compare_report()

    print("\n[3/4] 重复执行（判重）")
    again = run(DailyMergeOptions(day_dir=day_dir))
    check(again["results"][KIND_XLS]["added"] == 0, "xls 第二次新增 0 行")
    check(again["results"][KIND_MDB]["added"] == 0, "mdb 第二次新增 0 行")
    check(again["results"][KIND_REPORT].get("skipped"), "报告第二次未重复粘贴")
    check(not again["backups"], "没有改动就不产生多余备份")

    print("\n[4/4] 还原备份")
    done = restore(os.path.join(SANDBOX, "北部山头3940平台2026"), "1008")
    if done:
        check(len(done) == 3, "三份总表都从备份还原")
        preview = run(DailyMergeOptions(day_dir=day_dir, dry_run=True))
        check(preview["results"][KIND_XLS]["added"] == 150, "还原后 xls 又是 70→120")
        check(preview["results"][KIND_MDB]["added"] == 150, "还原后 mdb 又是 70→120")
    else:
        print("   （本次没有产生备份——总表本来就没变过，跳过还原检查）")
        check(True, "无备份可还原（数据本已并入）")

    failed = [msg for ok, msg in _checks if not ok]
    print("\n" + "=" * 78)
    print("共 %d 项检查，%d 项通过" % (len(_checks), len(_checks) - len(failed)))
    if failed:
        for msg in failed:
            print("  FAIL  %s" % msg)
        print("结论：存在差异 ✗")
        return 1
    print("结论：全部一致 ✓")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
