# -*- coding: utf-8 -*-
"""
3DMine 二次圈矿工具 —— 原型 V2 主程序（多品位等级 Voronoi 圈矿）
==================================================================
核心功能（按用户要求）：
  1. 分别自动输出 4 个品位等级的二次圈矿矿界线文件（.str）：
       0.5-1 g/t、1-1.5 g/t、1.5-3 g/t、≥3 g/t
  2. 通过赋 Voronoi（泰森多边形）图计算各矿块矿量与品位
  3. 给每个矿块添加边长为 1 m 的正方形网格图
  4. 全部结果输出到 CAD（DXF 文件，按等级分图层）

运行方式：
    python main.py [--config config.json] [--csv sample_data/blast_BL2026-015.csv]

算法模式：--method voronoi（默认，用户指定）/ idw（计划书 IDW+等值线，保留）
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "src"
sys.path.insert(0, str(SRC))

from data_loader import prepare_hole_points                 # noqa: E402
from dxf_writer import write_dxf                            # noqa: E402
from grid import grid_lines_inside                         # noqa: E402
from str_writer import write_str                           # noqa: E402
from voronoi import (GRADE_LEVELS, OreBlock, absorb_low_grade_cells,  # noqa: E402
                     build_voronoi_cells, classify_cells, merge_to_blocks)

# ---------- 默认参数（计划书 14.1 金矿推荐值） ----------
DEFAULT_CONFIG = {
    "blast_id": "BL-2026-015",
    "platform": "3940",          # 平台（工程号前缀解析，如 BS-3940-0915 → 3940）
    "date": "0915",              # 日期（如 0915）
    "bench_id": "1220",
    "bench_z": 1220.0,
    "bench_height_m": 6.0,        # 台阶高度（mdb 真实数据默认 10：孔深 11 = 台阶 10 + 超钻 1）
    "top_cut_pct": 98.0,          # 特高品位切割分位
    "grid_res_m": 1.0,            # 矿块网格边长（1 m）
    "min_ore_area_m2": 15.0,      # 最小可采块体面积（真实孔距密，单元约 20~44 m²/孔）
    "density_ore_t_m3": 2.7,      # 矿石容重（真实示例：12590.453/4663.131 = 2.7 t/m³）
    "method": "voronoi",          # voronoi（默认）| idw
    # 矿块品位归并（用户规则：每矿块单一品位，低品位孔并入邻接高档块）
    "absorb_low_grade_cells": True,   # 启用跨档品位归并
    "absorb_max_low_ratio": 0.5,      # 低档面积 ≤ 高档面积的该比例才并入
    # IDW 模式参数（保留）
    "idw_cutoff_grade": 1.0,
    "idw_power": 2.0,
    "search_radius_m": 12.0,
    "min_waste_area_m2": 200.0,
    "simplify_tol_m": 0.5,
    "min_edge_len_m": 3.0,
}

GRADE_BIN_LABEL = {lv["id"]: lv["label"] for lv in GRADE_LEVELS}
# 报告类型标签（对齐人工报告：3 位小数区间，如 1.500-3.000）
GRADE_REPORT_LABEL = {
    "L1": "0.500-1.000",
    "L2": "1.000-1.500",
    "L3": "1.500-3.000",
    "L4": "3.000-999.000",
}
# 体号编号顺序：与人工报告一致，高档（≥3）在前 → 低档（0.5-1）在后
BLOCK_ORDER = ["L4", "L3", "L2", "L1"]


def run_voronoi(config: dict, comp: pd.DataFrame, out_dir: Path) -> dict:
    """Voronoi 分级圈矿主流程（输入为已合成的台阶品位点表）。"""
    print("=" * 64)
    print(f"爆区 {config['blast_id']} / 台阶 {config['bench_id']} / Voronoi 分级圈矿")
    print("=" * 64)
    print(f"[数据] 可用炮孔 {len(comp)} 个")

    # 1) Voronoi 构建与品位分级
    cells = build_voronoi_cells(
        comp["X"].to_numpy(float), comp["Y"].to_numpy(float),
        comp["Grade"].to_numpy(float), comp["Hole_ID"].tolist(),
    )
    grouped = classify_cells(cells)
    n_waste = len(cells) - sum(len(v) for v in grouped.values())
    for lv in GRADE_LEVELS:
        print(f"[分级] {lv['label']}: {len(grouped[lv['id']])} 个泰森单元"
              + ("" if lv["id"] != "L1" else f"，废石单元 {n_waste} 个"))

    # 2) 各等级合并为候选矿块并计算矿量/品位
    blocks_by_level = {}
    for lv in GRADE_LEVELS:
        blocks_by_level[lv["id"]] = merge_to_blocks(
            grouped[lv["id"]],
            bench_height_m=config["bench_height_m"],
            density_ore_t_m3=config["density_ore_t_m3"],
            min_area_m2=config["min_ore_area_m2"],
        )

    # 3) 矿块品位归并：低品位孔并入邻接高档块（每矿块单一品位）
    if config.get("absorb_low_grade_cells", True):
        blocks_by_level = absorb_low_grade_cells(
            blocks_by_level,
            bench_height_m=config["bench_height_m"],
            density_ore_t_m3=config["density_ore_t_m3"],
            max_low_ratio=config.get("absorb_max_low_ratio", 0.5),
        )

    # 4) 全局矿块编号（与人工报告一致：高档→低档，L4/L3/L2/L1）
    all_blocks: list[OreBlock] = []
    for lv in BLOCK_ORDER:
        for b in blocks_by_level.get(lv, []):
            b.block_no = len(all_blocks) + 1
            all_blocks.append(b)
            print(f"[矿块] {GRADE_BIN_LABEL[lv]} 体号#{b.block_no}: "
                  f"{b.area_m2:.0f} m² / 体积 {b.volume_m3:.0f} m³ / 重量 {b.tonnage_t:.1f} t"
                  f" / 品位 {b.avg_grade_g_t:.3f} g/t（{b.cell_count} 孔，含归并）")

    # 5) 分等级 .str 输出（每个等级一个矿界线文件）
    str_files = {}
    for lv in GRADE_LEVELS:
        blocks = [b for b in all_blocks if b.level_id == lv["id"]]
        if not blocks:
            continue
        polys = [np.asarray(b.polygon.exterior.coords) for b in blocks]
        fname = f"{config['blast_id']}_{lv['id']}_grade.str"
        write_str(polys, out_dir / fname, bench_z=config["bench_z"], code=lv["code"])
        str_files[lv["id"]] = fname
        print(f"[输出] .str：{fname}（{len(polys)} 条闭合界线）")

    # 6) 每个矿块 1m 正方形网格
    grid_lines: dict[int, list] = {}
    for b in all_blocks:
        gl = grid_lines_inside(b.polygon, cell_size=config["grid_res_m"])
        grid_lines[b.block_no] = gl
        print(f"[网格] 体号#{b.block_no}（{b.level_id}）生成 1m 网格线段 {len(gl)} 条")

    # 7) CAD DXF 输出（含矿块号标注，文件名对齐人工"矿块图"）
    dxf_blocks = [{
        "level_id": b.level_id,
        "block_no": b.block_no,
        "polygon": b.polygon,
        "label": f"{b.block_no}",
    } for b in all_blocks]
    dxf_path = write_dxf(
        out_dir / f"{config['platform']}平台 {config['date']}二次圈矿矿块图.dxf",
        blocks=dxf_blocks,
        grid_lines=grid_lines,
        title=f"{config['platform']}平台 {config['date']} 二次圈矿矿块图（体号=品位归并后矿块号）",
    )
    print(f"[输出] CAD DXF：{dxf_path.name}")

    # 8) 炮孔数据报告（对齐人工格式：类型|体号|体积|重量|平均品位|金属量（百克））
    #    类型降序（≥3 → 1.5-3 → 1-1.5 → 0.5-1），每类型小计 + 合计；金属量单位百克（÷100）
    def report_rows():
        out_rows = []
        for lv in BLOCK_ORDER:
            bs = [b for b in all_blocks if b.level_id == lv]
            if not bs:
                continue
            for b in bs:
                out_rows.append({
                    "类型": GRADE_REPORT_LABEL[lv], "体号": b.block_no,
                    "体积": round(b.volume_m3, 3), "重量（t）": round(b.tonnage_t, 3),
                    "平均品位(Au)": round(b.avg_grade_g_t, 3),
                    "金属量（百克）": round(b.metal_g / 100.0, 3),
                })
            st = sum(b.tonnage_t for b in bs)
            sm = sum(b.metal_g for b in bs)
            out_rows.append({
                "类型": "小计", "体号": None,
                "体积": round(sum(b.volume_m3 for b in bs), 3), "重量（t）": round(st, 3),
                "平均品位(Au)": round(sm / st, 3) if st > 0 else None,
                "金属量（百克）": round(sm / 100.0, 3),
            })
        t_all = sum(b.tonnage_t for b in all_blocks)
        m_all = sum(b.metal_g for b in all_blocks)
        out_rows.append({
            "类型": "合计", "体号": None,
            "体积": round(sum(b.volume_m3 for b in all_blocks), 3), "重量（t）": round(t_all, 3),
            "平均品位(Au)": round(m_all / t_all, 3) if t_all > 0 else None,
            "金属量（百克）": round(m_all / 100.0, 3),
        })
        return out_rows

    report = report_rows()
    report_csv = out_dir / f"{config['platform']}平台 {config['date']}炮孔数据报告.csv"
    pd.DataFrame(report).to_csv(report_csv, index=False, encoding="utf-8-sig")

    # 9) 炮孔数据报告 xlsx（与人工 Excel 报告同版式）
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append([f"{config['platform']}平台 {config['date']}炮孔数据报告"])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append(["类型", "体号", "体积", "重量（t）", "平均品位(Au)", "金属量（百克）"])
        for r in report:
            ws.append([r["类型"], r["体号"], r["体积"], r["重量（t）"],
                       r["平均品位(Au)"], r["金属量（百克）"]])
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(horizontal="center")
        report_xlsx = out_dir / f"{config['platform']}平台 {config['date']}炮孔数据报告.xlsx"
        wb.save(str(report_xlsx))
    except Exception as e:
        report_xlsx = report_csv
        print(f"  [警告] xlsx 报告生成失败，已改输出 CSV：{e}")

    # 10) 矿块-炮孔明细表（可粘贴 CAD）+ 矿块汇总文本
    hole_rows = []
    for b in all_blocks:
        for c in b.cells:
            hole_rows.append({
                "体号": b.block_no, "类型": GRADE_REPORT_LABEL.get(b.level_id, ""),
                "Hole_ID": c.hole_id, "品位_g_t": round(c.grade, 3),
                "单元面积_m2": round(c.polygon.area, 1),
            })
    pd.DataFrame(hole_rows).to_csv(out_dir / f"{config['platform']}平台 {config['date']}矿块炮孔明细.csv",
                                   index=False, encoding="utf-8-sig")

    txt_lines = [f"{'体号':<6}{'品位类型':<14}{'面积m2':>8}{'体积m3':>10}"
                 f"{'重量t':>10}{'平均品位g/t':>12}{'金属量g':>10}  炮孔数  炮孔号"]
    for b in all_blocks:
        holes_s = "、".join(c.hole_id.split("-")[-1] for c in b.cells)
        txt_lines.append(
            f"#{b.block_no:<5}{GRADE_REPORT_LABEL.get(b.level_id, ''):<14}{b.area_m2:>8.0f}"
            f"{b.volume_m3:>10.0f}{b.tonnage_t:>10.0f}{b.avg_grade_g_t:>12.3f}{b.metal_g:>10.0f}"
            f"  {b.cell_count:<6}{holes_s}")
    txt_path = out_dir / f"{config['platform']}平台 {config['date']}矿块汇总（粘贴CAD）.txt"
    txt_path.write_text("\n".join(txt_lines), encoding="utf-8")

    # 等级汇总
    level_summary = {}
    for lv in GRADE_LEVELS:
        bs = [b for b in all_blocks if b.level_id == lv["id"]]
        if bs:
            t = sum(b.tonnage_t for b in bs)
            m = sum(b.metal_g for b in bs)
            level_summary[lv["id"]] = {
                "grade_range": lv["label"], "block_count": len(bs),
                "tonnage_t": round(t, 1), "metal_g": round(m, 0),
                "avg_grade_g_t": round(m / t, 3) if t > 0 else 0.0,
            }

    result = {
        "blast_id": config["blast_id"],
        "platform": config["platform"], "date": config["date"],
        "bench_id": config["bench_id"],
        "method": "voronoi",
        "config": config,
        "absorb_enabled": config.get("absorb_low_grade_cells", True),
        "level_summary": level_summary,
        "blocks": report,
        "block_holes": hole_rows,
        "str_files": str_files,
        "dxf_file": dxf_path.name,
        "report_xlsx": report_xlsx.name,
    }
    (out_dir / f"{config['blast_id']}_result.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"[报告] 炮孔数据报告（xlsx+csv）/矿块炮孔明细/汇总文本/结果 JSON 已写出至 {out_dir}")

    return {"all_blocks": all_blocks, "cells": cells, "grouped": grouped,
            "grid_lines": grid_lines, "comp": comp, "result": result,
            "blocks_by_level": blocks_by_level}


def plot_validation(res: dict, config: dict, out_dir: Path) -> Path:
    """验证图：Voronoi 单元按等级着色 + 各等级矿块边界 + 1m 网格（示意主矿块）。"""
    from matplotlib import pyplot as plt
    from matplotlib.patches import Polygon as MplPolygon

    plt.rcParams["font.sans-serif"] = ["Microsoft YaHei", "SimHei"]
    plt.rcParams["axes.unicode_minus"] = False

    # 颜色对齐 3DMine style_table_3dmine：0.5-1 青 / 1-1.5 蓝 / 1.5-3 红 / ≥3 黄
    color_map = {"L1": "cyan", "L2": "blue", "L3": "red", "L4": "yellow"}
    label_map = GRADE_BIN_LABEL

    fig, ax = plt.subplots(figsize=(12, 9.5))
    for c in res["cells"]:
        if c.level_id is None:
            continue
        color = color_map[c.level_id]
        poly = MplPolygon(list(c.polygon.exterior.coords), closed=True,
                          facecolor=color, edgecolor="none", alpha=0.55)
        ax.add_patch(poly)

    # 等级边界 + 矿块号标注
    for b in res["all_blocks"]:
        ext = np.asarray(b.polygon.exterior.coords)
        ax.plot(ext[:, 0], ext[:, 1], color="black", lw=1.8)
        cx, cy = b.polygon.representative_point().coords[0]
        ax.text(cx, cy, f"#{b.block_no}", fontsize=9, ha="center", va="center",
                color="black", fontweight="bold",
                bbox=dict(facecolor="white", alpha=0.8, edgecolor="none", pad=1.0))

    # 网格（主矿块示意，仅画第一个矿块的 1m 网格）
    blocks = res["all_blocks"]
    if blocks:
        b0 = blocks[0]
        gl = grid_lines_inside(b0.polygon, cell_size=config["grid_res_m"])
        for ln in gl:
            xs, ys = zip(*ln.coords)
            ax.plot(xs, ys, color="gray", lw=0.4, alpha=0.8)

    # 图例
    handles = [plt.Line2D([0], [0], marker="s", color="w", markerfacecolor=c, markersize=12,
                          label=label_map[lv]) for lv, c in color_map.items()]
    handles.append(plt.Line2D([0], [0], color="gray", lw=0.6, label="1m 网格（主矿块示意）"))
    ax.legend(handles=handles, loc="upper left", fontsize=9)

    ax.set_title(f"爆区 {config['blast_id']} · 台阶 {config['bench_id']} · Voronoi 分级圈矿\n"
                 f"4 个品位等级分别输出矿界线 · 矿块叠加 1m 网格", fontsize=13)
    ax.set_xlabel("X (m)")
    ax.set_ylabel("Y (m)")
    ax.set_aspect("equal")
    fig.tight_layout()

    png_path = out_dir / f"{config['platform']}平台 {config['date']}二次圈矿验证图.png"
    fig.savefig(png_path, dpi=150)
    plt.close(fig)
    print(f"[验证图] 已生成：{png_path}")
    return png_path


def run_idw(config: dict, csv_path: Path, out_dir: Path) -> dict:
    """IDW + 等值线模式（保留计划书路径 A，便于对比；单一边界品位）。"""
    from interpolation import (apply_mining_rules, build_grid,
                               extract_ore_boundaries, idw_interpolate)
    from mining_rules import filter_by_area, straighten_edges
    from reserve import reserve_for_polygons

    print("=" * 64)
    print(f"爆区 {config['blast_id']} / IDW 等值线模式（边界品位 {config['idw_cutoff_grade']} g/t）")
    print("=" * 64)
    comp, meta = prepare_hole_points(csv_path, cutoff_top_pct=config["top_cut_pct"])
    x = comp["X"].to_numpy(float); y = comp["Y"].to_numpy(float)
    g = comp["Grade"].to_numpy(float)
    xg, yg = build_grid(x, y, config["grid_res_m"])
    grid_z = idw_interpolate(x, y, g, xg, yg, power=config["idw_power"],
                             search_radius=config["search_radius_m"])
    grid_z = apply_mining_rules(grid_z, config["idw_cutoff_grade"],
                                min_ore_area_m2=config["min_ore_area_m2"],
                                min_waste_area_m2=config["min_waste_area_m2"],
                                resolution=config["grid_res_m"])
    polys = extract_ore_boundaries(grid_z, xg, yg, config["idw_cutoff_grade"],
                                   config["simplify_tol_m"])
    polys = [straighten_edges(p, config["min_edge_len_m"]) for p in polys]
    polys = filter_by_area(polys, config["min_ore_area_m2"])
    str_path = out_dir / f"{config['blast_id']}_ore_boundary.str"
    write_str(polys, str_path, bench_z=config["bench_z"])
    reserve = reserve_for_polygons(polys, grid_z, xg, yg,
                                   config["bench_height_m"],
                                   config["density_ore_t_m3"],
                                   config["idw_cutoff_grade"])
    s = reserve["summary"]
    print(f"[IDW] 矿块 {len(polys)} 个，矿石量 {s['Total_Tonnage_t']} t，"
          f"平均品位 {s['Weighted_Avg_Grade_g_t']} g/t，金属量 {s['Total_Metal_g']} g")
    print("（IDW 模式输出单一边界品位，多等级请使用默认 voronoi 模式）")
    return {"reserve": reserve}


def main() -> None:
    ap = argparse.ArgumentParser(description="3DMine 二次圈矿工具原型 V2（多等级 Voronoi）")
    ap.add_argument("--config", type=str, default="", help="参数配置文件路径（JSON）")
    ap.add_argument("--csv", type=str, default="", help="炮孔化验 CSV 路径")
    ap.add_argument("--method", type=str, choices=["voronoi", "idw"], default="",
                    help="圈矿算法：voronoi（默认）| idw")
    args = ap.parse_args()

    config = dict(DEFAULT_CONFIG)
    if args.config:
        cfg_path = Path(args.config)
        if not cfg_path.exists():
            print(f"配置不存在：{cfg_path}")
            sys.exit(1)
        config.update(json.loads(cfg_path.read_text(encoding="utf-8")))
    if args.method:
        config["method"] = args.method

    csv_path = Path(args.csv) if args.csv else ROOT / "sample_data" / "blast_BL2026-015.csv"
    if not csv_path.exists():
        print(f"数据文件不存在：{csv_path}")
        sys.exit(1)

    out_dir = ROOT / "output"
    out_dir.mkdir(parents=True, exist_ok=True)

    if config["method"] != "voronoi":
        run_idw(config, csv_path, out_dir)
        print("\n全流程完成。")
        return

    # Voronoi 模式：数据源支持 3DMine mdb / Excel xls / CSV
    #   mdb/xls：定位表 + 测斜表（直孔）+ pit_blast 品位表（人工流程的权威数据源）
    #   CSV：模拟数据或按同字段导出的文本
    if csv_path.suffix.lower() == ".mdb":
        from mdb_loader import load_mdb_blast
        comp_all = load_mdb_blast(csv_path)
        print(f"[数据] mdb 读取完成：{len(comp_all)} 个炮孔（{csv_path.name}）")
        group_cols = ["Platform_ID", "Date"]
        # 真实矿山台阶：孔深 10~11 m = 台阶 10 m + 超钻（config.json 可显式覆盖）
        real_bench_height = 10.0
    elif csv_path.suffix.lower() in (".xls", ".xlsx"):
        from mdb_loader import load_xls_blast
        comp_all = load_xls_blast(csv_path)
        print(f"[数据] xls 读取完成：{len(comp_all)} 个炮孔（{csv_path.name}）")
        group_cols = ["Platform_ID", "Date"]
        real_bench_height = 10.0
    else:
        comp_all, meta = prepare_hole_points(csv_path, cutoff_top_pct=None)
        print(f"[数据] CSV 读取完成：{meta['usable_holes']} 个炮孔（合成品位点 {len(comp_all)} 条）")
        group_cols = [c for c in ("Platform_ID", "Date") if c in comp_all.columns] or ["Blast_ID"]

    n_groups = 0
    for key, grp in comp_all.groupby(group_cols, sort=False):
        key_list = key if isinstance(key, tuple) else (key,)
        platform, date = str(key_list[0]), str(key_list[1]) if len(key_list) > 1 else str(key_list[0])
        cfg = dict(config)
        cfg["platform"] = platform
        cfg["date"] = date
        cfg["blast_id"] = f"BS-{platform}-{date}" if date else platform
        if csv_path.suffix.lower() in (".mdb", ".xls", ".xlsx"):
            cfg["bench_height_m"] = real_bench_height
        if "Bench_ID" in grp.columns and grp["Bench_ID"].nunique() == 1:
            cfg["bench_id"] = str(grp["Bench_ID"].iloc[0])
        if grp["Z"].nunique() > 0:
            cfg["bench_z"] = float(grp["Z"].median())
        res = run_voronoi(cfg, grp.reset_index(drop=True), out_dir)
        plot_validation(res, cfg, out_dir)
        n_groups += 1
    print(f"\n全流程完成，共圈连 {n_groups} 组（平台/日期）。")


if __name__ == "__main__":
    main()
