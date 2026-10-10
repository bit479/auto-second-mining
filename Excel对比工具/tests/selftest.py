"""端到端自测：造两份带已知差异的样例，跑对比，校验报告内容。"""
import os
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

import openpyxl
from compare import compare_workbooks, diff_count
from reader import read_workbook
from writer import write_report

tmp = tempfile.mkdtemp(prefix="excel_diff_test_")

# ---- 构造样例 ----
def build_a(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["姓名", "年龄", "城市"])
    ws.append(["张三", 20, "北京"])
    ws.append(["李四", 25, "上海"])
    ws.append(["王五", 30, "广州"])
    cfg = wb.create_sheet("配置")
    cfg.append(["参数", "值"])
    cfg.append(["超时", 30])
    cfg.append(["重试", 3])
    wb.save(path)

def build_b(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "数据"
    ws.append(["姓名", "年龄", "城市"])
    ws.append(["张三", 21, "北京"])      # 年龄改了
    ws.append(["李四", 25, "深圳"])      # 城市改了
    ws.append(["王五", 30, "广州"])
    ws.append(["赵六", 35, "杭州"])      # 仅B有
    wb.save(path)

pa = os.path.join(tmp, "A.xlsx")
pb = os.path.join(tmp, "B.xlsx")
out = os.path.join(tmp, "报告.xlsx")
build_a(pa)
build_b(pb)

wb_a = read_workbook(pa)
wb_b = read_workbook(pb)
result = compare_workbooks(wb_a, wb_b)
write_report(result, pa, pb, out, source_a=wb_a, source_b=wb_b)

t = result["totals"]
n = diff_count(result)
print("totals:", t, "diff_total:", n)

# ---- 校验 ----
# 配置 sheet 非空格: 参数/值/超时/30/重试/3 = 6 个 -> only_a 应为 6
errs = []
if t["changed"] != 2: errs.append(f"changed 应为2，实际 {t['changed']}")
if t["only_a"] != 6: errs.append(f"only_a 应为6，实际 {t['only_a']}")
if t["only_b"] != 3: errs.append(f"only_b 应为3，实际 {t['only_b']}")
if n != 11: errs.append(f"diff_total 应为11，实际 {n}")

# 读回报告校验 sheet 与差异清单行数
rep = openpyxl.load_workbook(out)
print("sheets:", rep.sheetnames)
need = {"汇总", "差异清单", "对比_数据", "对比_配置"}
if not need.issubset(set(rep.sheetnames)):
    errs.append(f"缺少 sheet: {need - set(rep.sheetnames)}")

dl = rep["差异清单"]
rows = dl.max_row - 1  # 去掉表头
if rows != 11:
    errs.append(f"差异清单应为11行，实际 {rows}")

# 校验高亮：对比_数据 中 B2(年龄, 第2行第2列) 应黄，第3行第3列(城市) 应黄，第5行应红
vis = rep["对比_数据"]
fc = vis.cell(row=2, column=2).fill.fgColor.rgb
if fc != "00FFF2CC":
    errs.append(f"数据!B2 高亮异常: {fc}")
fc2 = vis.cell(row=3, column=3).fill.fgColor.rgb
if fc2 != "00FFF2CC":
    errs.append(f"数据!C3 高亮异常: {fc2}")
# 第5行(赵六)整行应为红
red = vis.cell(row=5, column=1).fill.fgColor.rgb
if red != "00FFC7CE":
    errs.append(f"数据!A5 高亮异常: {red}")
# 未差异的 A1 应无填充 (00000000)
none_fill = vis.cell(row=1, column=1).fill.fgColor.rgb
if none_fill not in (None, "00000000"):
    errs.append(f"数据!A1 不应高亮: {none_fill}")

if errs:
    print("❌ 测试未通过:")
    for e in errs:
        print("  -", e)
    sys.exit(1)
else:
    print("✅ 全部通过：差异计数、sheet 结构、高亮均正确")
    print(f"报告路径: {out}")
    print(f"样例目录: {tmp}")
